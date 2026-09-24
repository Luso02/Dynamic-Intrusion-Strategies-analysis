#!/usr/bin/env python3
"""
Tunes the SAHybrid attack strategy with Optuna, then compares it against the
MaxProbability and Stealth baselines using parallel runs and confidence intervals.
"""

import csv
import json
import os
import asyncio
import aiohttp
import nest_asyncio
import random
import math
import numpy as np
import re
import ipaddress
import traceback
import multiprocessing
from collections import defaultdict
from datetime import datetime
from concurrent.futures import ProcessPoolExecutor, as_completed
from scipy import stats as sp_stats

try:
    import optuna
    from optuna.pruners import MedianPruner
    from optuna.samplers import TPESampler
    HAS_OPTUNA = True
except ImportError:
    HAS_OPTUNA = False

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

import networkx as nx


# Confidence intervals

def compute_ci(samples, confidence=0.95):
    """Mean and CI half-width (Student's t). Returns (mean, ci, std)."""
    samples = np.array(samples, dtype=float)
    samples = samples[~np.isnan(samples)]
    n = len(samples)
    if n == 0:
        return (np.nan, np.nan, np.nan)
    if n == 1:
        return (float(samples[0]), 0.0, 0.0)
    mean = np.mean(samples)
    std = np.std(samples, ddof=1)
    se = std / np.sqrt(n)
    t_crit = sp_stats.t.ppf((1 + confidence) / 2, df=n - 1)
    ci = t_crit * se
    return (float(mean), float(ci), float(std))


# Assets and digital twin

class Asset:
    def __init__(self, name, **kwargs):
        self.name = name
        self.attributes = kwargs
    def __repr__(self):
        return f"{self.__class__.__name__}(name='{self.name}')"

class Host(Asset):
    pass

class Software(Asset):
    def __init__(self, name, **kwargs):
        super().__init__(name, **kwargs)
        self.base_name = kwargs.get('base_name')
        self.version = kwargs.get('version')

class VirtualMachine(Host):
    pass


class CVEEnricher:
    def __init__(self, graph, api_key=None, cache_file='cve_cache.json'):
        self.graph = graph
        self.api_key = api_key
        self.cache_file = cache_file
        self.cve_cache = self._load_cache()

    def _load_cache(self):
        if os.path.exists(self.cache_file):
            with open(self.cache_file, 'r') as f:
                return json.load(f)
        return {}

    def _save_cache(self):
        with open(self.cache_file, 'w') as f:
            json.dump(self.cve_cache, f, indent=2)

    def _normalize_sw_name(self, name):
        return name.split('-')[0]

    async def _fetch_cve(self, session, sw_name, clean_version, original_key):
        if original_key in self.cve_cache:
            return original_key, self.cve_cache[original_key], '.'
        search_name = self._normalize_sw_name(sw_name)
        headers = {'apiKey': self.api_key} if self.api_key else {}
        query = f"keywordSearch={search_name} {clean_version}"
        url = f"https://services.nvd.nist.gov/rest/json/cves/2.0?{query}"
        try:
            async with session.get(url, headers=headers, timeout=40) as response:
                response.raise_for_status()
                data = await response.json()
                vulnerabilities = []
                if 'vulnerabilities' in data:
                    for item in data['vulnerabilities']:
                        cve = item['cve']
                        score = -1.0
                        if 'cvssMetricV31' in cve['metrics']:
                            score = cve['metrics']['cvssMetricV31'][0]['cvssData']['baseScore']
                        elif 'cvssMetricV2' in cve['metrics']:
                            score = cve['metrics']['cvssMetricV2'][0]['cvssData']['baseScore']
                        vulnerabilities.append({'id': cve['id'], 'score': score})
                self.cve_cache[original_key] = vulnerabilities
                feedback_char = 'V' if vulnerabilities else '.'
                return original_key, vulnerabilities, feedback_char
        except Exception:
            self.cve_cache[original_key] = []
            return original_key, [], 'E'

    async def run_enrichment(self):
        unique_software = {
            f"{data.get('base_name', '').lower()}|{data.get('version', '')}": (
                data.get('base_name'),
                re.match(r'[\d\.:]+', data.get('version', '')).group(0)
                if re.match(r'[\d\.:]+', data.get('version', ''))
                else data.get('version', '')
            )
            for _, data in self.graph.nodes(data=True)
            if data.get('type') == 'Software' and data.get('base_name') and data.get('version')
        }
        items_to_fetch = [
            (key, sw_name, clean_version)
            for key, (sw_name, clean_version) in unique_software.items()
            if key not in self.cve_cache
        ]
        if items_to_fetch:
            async with aiohttp.ClientSession() as session:
                await asyncio.gather(
                    *[self._fetch_cve(session, sw, ver, key) for key, sw, ver in items_to_fetch]
                )
        print(f"CVE: {len(unique_software)} software, {len(items_to_fetch)} fetch eseguiti.")
        for node_name, data in self.graph.nodes(data=True):
            if data.get('type') == 'Software':
                key = f"{data.get('base_name', '').lower()}|{data.get('version', '')}"
                if (vulns := self.cve_cache.get(key)):
                    self.graph.nodes[node_name]['vulnerabilities'] = vulns
                    self.graph.nodes[node_name]['max_cvss'] = max(
                        [v['score'] for v in vulns if v['score'] >= 0] or [0.0]
                    )
        self._save_cache()


class DigitalTwin:
    def __init__(self):
        self.graph = nx.DiGraph()
        self.assets = {}
        self.discovered_subnets = []

    def _parse_multiline_data(self, data: str):
        return [item.strip() for item in data.split('<br>') if item.strip()] if isinstance(data, str) else []

    def _parse_software_data(self, data: str):
        return [
            {'name': parts[0].strip(), 'version': parts[1].strip()}
            for line in self._parse_multiline_data(data)
            if len(parts := line.rsplit(' - ', 1)) == 2
        ]

    def _discover_subnets_from_rows(self, all_rows):
        subnet_prefixes = set()
        for row in all_rows:
            for ip_str in self._parse_multiline_data(row.get('Networking - IP', '')):
                try:
                    ip_addr = ipaddress.ip_address(ip_str)
                    if not (ip_addr.is_loopback or ip_addr.version == 6 or not ip_addr.is_private):
                        subnet_prefixes.add('.'.join(ip_str.split('.')[:3]))
                except ValueError:
                    continue
        self.discovered_subnets = [
            ipaddress.ip_network(f"{prefix}.0/24") for prefix in sorted(subnet_prefixes)
        ]

    def _get_subnet_for_ips(self, ips):
        for ip_str in ips:
            try:
                subnet = next(
                    (str(s) for s in self.discovered_subnets if ipaddress.ip_address(ip_str) in s), None
                )
                if subnet:
                    return subnet
            except ValueError:
                continue
        return 'Unknown/External'

    def _add_or_get_asset(self, name, asset_class, **kwargs):
        if name not in self.assets:
            self.assets[name] = asset_class(name, **kwargs)
            self.graph.add_node(name, type=asset_class.__name__, **kwargs)
        elif kwargs:
            nx.set_node_attributes(self.graph, {name: kwargs})
        return self.assets[name]

    def load_from_csv(self, file_path, delimiter=';'):
        with open(file_path, mode='r', encoding='utf-8-sig') as infile:
            next(infile)
            all_rows = list(csv.DictReader(infile, delimiter=delimiter))
        self._discover_subnets_from_rows(all_rows)
        for row in all_rows:
            if not (host_name := row.get('Name')):
                continue
            ips = self._parse_multiline_data(row.get('Networking - IP', ''))
            subnet = self._get_subnet_for_ips(ips)
            host_attributes = {
                'os': row.get('Operating System - Name'),
                'asset_type': row.get('Type'), 'ips': ips, 'subnet': subnet
            }
            self._add_or_get_asset(
                host_name, VirtualMachine if host_attributes['asset_type'] == 'VM' else Host,
                **host_attributes
            )
            for vm_name in self._parse_multiline_data(row.get('Virtual machines - Name', '')):
                self._add_or_get_asset(vm_name, VirtualMachine, os=host_attributes['os'], subnet=subnet)
                self.graph.add_edge(host_name, vm_name, relationship='HOSTS')
            for sw in self._parse_software_data(row.get('Software - Name', '')):
                sw_unique_name = f"{sw['name']} ({sw['version']})"
                self._add_or_get_asset(sw_unique_name, Software, base_name=sw['name'], version=sw['version'])
                self.graph.add_edge(host_name, sw_unique_name, relationship='INSTALLS')
        print(f"Grafo costruito: {self.graph.number_of_nodes()} nodi, {self.graph.number_of_edges()} archi.")

    def get_graph(self):
        return self.graph


# Attacker state and strategies

class AttackerState:
    def __init__(self, initial_host):
        self.compromised_hosts = {initial_host: {'priv_level': 'user'}}
        self.credentials = set()
        self.current_time_hours = 0.0
        self.path = ['Attacker', initial_host]
        self.detected = False
        self.goal_achieved = False
        self.noise_level = 0.0


class AttackStrategy:
    def __init__(self, simulator):
        self.sim = simulator
    def choose_next_action(self, state, current_host):
        raise NotImplementedError
    def choose_target(self, state, current_host, available_targets):
        raise NotImplementedError
    def __str__(self):
        return self.__class__.__name__


class StealthStrategy(AttackStrategy):
    def __init__(self, simulator, noise_threshold=10.0):
        super().__init__(simulator)
        self.noise_threshold = noise_threshold
    def choose_next_action(self, state, host):
        if state.noise_level > self.noise_threshold:
            return 'WAIT'
        priv_level = state.compromised_hosts[host]['priv_level']
        if 'domain_admin' in state.credentials or 'local_admin' in state.credentials:
            return 'LATERAL_MOVE'
        if priv_level == 'user':
            return 'PRIV_ESC'
        if priv_level == 'admin':
            return 'DUMP_CREDS'
        return 'LATERAL_MOVE'
    def choose_target(self, state, current_host, targets):
        if not targets:
            return None
        return min(targets, key=lambda t: self.sim.dt_graph.nodes[t].get('value', 1))


class MaxProbabilityStrategy(AttackStrategy):
    def choose_next_action(self, state, host):
        if state.compromised_hosts[host]['priv_level'] == 'user':
            return 'PRIV_ESC'
        if 'domain_admin' not in state.credentials and self.sim.dt_graph.nodes[host].get('value', 0) >= 1000:
            return 'DUMP_CREDS'
        if 'local_admin' not in state.credentials:
            return 'DUMP_CREDS'
        return 'LATERAL_MOVE'
    def _get_max_cvss_for_host(self, target_host):
        vulns = [
            self.sim.dt_graph.nodes[s].get('max_cvss', 0.0)
            for s in self.sim.dt_graph.successors(target_host)
            if self.sim.dt_graph.nodes[s].get('type') == 'Software'
            and self.sim.dt_graph.nodes[s].get('vulnerabilities')
        ]
        return max(vulns) if vulns else 0.0
    def choose_target(self, state, current_host, targets):
        if not targets:
            return None
        return max(targets, key=lambda t: self._get_max_cvss_for_host(t))


class SimulatedAnnealingStrategy(AttackStrategy):
    def __init__(self, simulator, T0=0.0, delta_T=1.0, T_max=15.0, gamma=0.5,
                 L=2, p_func='linear', noise_threshold=10.0):
        super().__init__(simulator)
        self.T0 = T0
        self.delta_T = delta_T
        self.T_max = T_max
        self.gamma = gamma
        self.L = L
        self.p_func = p_func
        self.noise_threshold = noise_threshold
        self.stealth = StealthStrategy(simulator, noise_threshold=noise_threshold)
        self.maxprob = MaxProbabilityStrategy(simulator)
        self._total_activations = 0
        self._total_maxprob_steps = 0
        self._total_runs = 0

    def _ensure_sa_state(self, state):
        if not hasattr(state, '_sa_T'):
            self._total_runs += 1
            state._sa_T = self.T0
            state._sa_deviation_remaining = 0
            state._sa_current_mode = 'stealth'
            state._sa_last_hosts = len(state.compromised_hosts)
            state._sa_activations = 0

    def _p(self, T):
        if self.p_func == 'linear':
            return min(1.0, T / self.T_max)
        elif self.p_func == 'sigmoid':
            mid = self.T_max / 2
            k = 6.0 / self.T_max
            return 1.0 / (1.0 + math.exp(-k * (T - mid)))
        elif self.p_func == 'exponential':
            scale = self.T_max / 3.0
            return min(1.0, 1.0 - math.exp(-T / max(scale, 0.01)))
        return min(1.0, T / self.T_max)

    def choose_next_action(self, state, host):
        self._ensure_sa_state(state)
        if state._sa_deviation_remaining > 0:
            state._sa_deviation_remaining -= 1
            state._sa_current_mode = 'maxprob'
            self._total_maxprob_steps += 1
            if state._sa_deviation_remaining == 0:
                state._sa_T *= self.gamma
            return self.maxprob.choose_next_action(state, host)
        hosts_now = len(state.compromised_hosts)
        if not state.goal_achieved:
            state._sa_T += self.delta_T
        else:
            state._sa_T = max(0, state._sa_T - self.delta_T * 2)
        state._sa_last_hosts = hosts_now
        p = self._p(state._sa_T)
        if random.random() < p:
            deviation_len = random.randint(1, self.L)
            state._sa_deviation_remaining = deviation_len - 1
            state._sa_activations += 1
            state._sa_current_mode = 'maxprob'
            self._total_activations += 1
            self._total_maxprob_steps += 1
            if state._sa_deviation_remaining == 0:
                state._sa_T *= self.gamma
            return self.maxprob.choose_next_action(state, host)
        state._sa_current_mode = 'stealth'
        return self.stealth.choose_next_action(state, host)

    def choose_target(self, state, current_host, targets):
        self._ensure_sa_state(state)
        if state._sa_current_mode == 'maxprob':
            return self.maxprob.choose_target(state, current_host, targets)
        return self.stealth.choose_target(state, current_host, targets)

    def __str__(self):
        return f"SAHybrid({self.p_func},dT={self.delta_T},Tm={self.T_max},g={self.gamma},L={self.L})"


def create_sa_hybrid_strategy(T0=0.0, delta_T=1.0, T_max=15.0, gamma=0.5,
                               L=2, p_func='linear', noise_threshold=10.0):
    class ConfiguredSAHybrid(SimulatedAnnealingStrategy):
        __name__ = f"SAHybrid_{p_func}_dT{delta_T}_Tm{T_max}_g{gamma}_L{L}"
        def __init__(self, simulator):
            super().__init__(simulator, T0=T0, delta_T=delta_T, T_max=T_max,
                           gamma=gamma, L=L, p_func=p_func,
                           noise_threshold=noise_threshold)
    return ConfiguredSAHybrid


# Simulation engine

class AttackSimulator:
    TIME_MODEL = {
        'SCAN': (1.0, 1.0), 'EXPLOIT': (0.1, 2), 'PRIV_ESC': (0.2, 3),
        'DUMP_CREDS': (0.1, 0.5), 'LOGIN': (0.01, 0.1)
    }
    NOISE_MODEL = {'SCAN': 1.5, 'EXPLOIT': 2.0, 'PRIV_ESC': 1.0, 'DUMP_CREDS': 1.2, 'LOGIN': 0.2}

    def __init__(self, dt_graph, strategy_class, skill=0.8, detection_rate=0.05):
        self.dt_graph = dt_graph
        self.strategy = strategy_class(self)
        self.skill = skill
        self.detection_rate = detection_rate
        self.attack_graph = nx.DiGraph()
        self._assign_asset_values()
        self.stats = self._reset_stats()

    def _reset_stats(self):
        return {
            'total': 0, 'succeeded': 0, 'failed': 0, 'detected': 0,
            'actions': defaultdict(int), 'exploited_cves': defaultdict(int),
            'compromised_nodes': defaultdict(int),
            'successful_path_details': defaultdict(list),
            'detected_paths': []
        }

    def _assign_asset_values(self):
        value_map = {'dc': 1000, 'db': 500, 'sql': 500, 'fileserver': 200, 'admin': 100}
        for node, data in self.dt_graph.nodes(data=True):
            if data.get('type') in ['Host', 'VirtualMachine']:
                self.dt_graph.nodes[node]['value'] = max(
                    [1] + [v for keyword, v in value_map.items() if keyword in node.lower()]
                )

    def _select_entry_point(self):
        web_keywords = ['apache', 'nginx', 'iis', 'httpd', 'tomcat']
        candidates = [
            (host, 1 + data.get('value', 1) + 10 * any(
                any(key in sw_name for key in web_keywords)
                for sw in self.dt_graph.successors(host)
                if (sw_name := self.dt_graph.nodes[sw].get('base_name', '').lower())
            ))
            for host, data in self.dt_graph.nodes(data=True)
            if data.get('type') in ['Host', 'VirtualMachine'] and data.get('subnet') == 'Unknown/External'
        ]
        if not candidates:
            return random.choice(
                [n for n, d in self.dt_graph.nodes(data=True) if d.get('type') in ['Host', 'VirtualMachine']]
            )
        return random.choices(*zip(*candidates), k=1)[0]

    def _get_noise_for_action(self, action):
        base_noise = self.NOISE_MODEL.get(action, 1.0)
        multiplier = getattr(self.strategy, 'noise_multiplier', 1.0)
        return base_noise * multiplier

    def _run_single_intrusion(self):
        self.stats['total'] += 1
        entry_host = self._select_entry_point()
        sw, vuln = self._attempt_exploit_cve(entry_host)
        if not sw:
            self.stats['failed'] += 1
            return
        state = AttackerState(entry_host)
        state.current_time_hours += self._get_time('EXPLOIT')
        state.noise_level += self._get_noise_for_action('EXPLOIT')

        while not state.detected and not state.goal_achieved and state.current_time_hours < 120:
            host = state.path[-1]
            if self.dt_graph.nodes[host].get('value', 0) >= 1000:
                state.goal_achieved = True
                continue
            action = self.strategy.choose_next_action(state, host)
            if action == 'PRIV_ESC':
                self._attempt_privesc(state, host)
            elif action == 'DUMP_CREDS':
                self._attempt_credential_dump(state, host)
            elif action == 'LATERAL_MOVE':
                self._attempt_lateral_move(state, host)
            elif action == 'WAIT':
                state.current_time_hours += 12
                state.noise_level *= 0.75

        if state.goal_achieved or len(state.compromised_hosts) > 1:
            self.stats['succeeded'] += 1
            path_tuple = tuple(state.path)
            self.stats['successful_path_details'][path_tuple].append({
                'time': state.current_time_hours, 'steps': len(state.path)
            })
        elif state.detected:
            self.stats['detected'] += 1
            self.stats['detected_paths'].append(tuple(state.path))
        else:
            self.stats['failed'] += 1

    def _attempt_lateral_move(self, state, current_host):
        state.current_time_hours += self._get_time('SCAN')
        state.noise_level += self._get_noise_for_action('SCAN')
        if self._check_detection(state, 'SCAN'):
            return
        subnet = self.dt_graph.nodes[current_host].get('subnet')
        targets = [
            n for n, d in self.dt_graph.nodes(data=True)
            if d.get('type') in ['Host', 'VirtualMachine']
            and d.get('subnet') == subnet and n not in state.compromised_hosts
        ]
        target_host = self.strategy.choose_target(state, current_host, targets)
        if not target_host:
            return
        moved, sw, vuln = False, None, None
        if 'domain_admin' in state.credentials and random.random() < 0.95:
            moved = True
        elif 'local_admin' in state.credentials and random.random() < 0.6:
            moved = True
        if moved:
            state.current_time_hours += self._get_time('LOGIN')
            state.noise_level += self._get_noise_for_action('LOGIN')
        else:
            state.current_time_hours += self._get_time('EXPLOIT')
            state.noise_level += self._get_noise_for_action('EXPLOIT')
            if self._check_detection(state, 'EXPLOIT'):
                return
            sw, vuln = self._attempt_exploit_cve(target_host)
            if sw:
                moved = True
        if moved:
            state.compromised_hosts[target_host] = {'priv_level': 'user'}
            state.path.append(target_host)

    def _get_time(self, action):
        return random.uniform(*self.TIME_MODEL[action]) / self.skill

    def _check_detection(self, state, action):
        if random.random() < self.detection_rate * self._get_noise_for_action(action):
            state.detected = True
            return True
        return False

    def _attempt_exploit_cve(self, target_host):
        vulns = [
            (s, max(self.dt_graph.nodes[s].get('vulnerabilities', []), key=lambda v: v['score']))
            for s in self.dt_graph.successors(target_host)
            if self.dt_graph.nodes[s].get('type') == 'Software'
            and self.dt_graph.nodes[s].get('vulnerabilities')
        ]
        if not vulns:
            return None, None
        sw_node, best_vuln = max(vulns, key=lambda item: item[1]['score'])
        if random.random() < (best_vuln['score'] / 10.0) * self.skill:
            return sw_node, best_vuln
        return None, None

    def _attempt_privesc(self, state, host):
        state.current_time_hours += self._get_time('PRIV_ESC')
        state.noise_level += self._get_noise_for_action('PRIV_ESC')
        if self._check_detection(state, 'PRIV_ESC'):
            return
        if random.random() < 0.7 * self.skill:
            state.compromised_hosts[host]['priv_level'] = 'admin'

    def _attempt_credential_dump(self, state, host):
        if state.compromised_hosts[host]['priv_level'] != 'admin':
            return
        state.current_time_hours += self._get_time('DUMP_CREDS')
        state.noise_level += self._get_noise_for_action('DUMP_CREDS')
        if self._check_detection(state, 'DUMP_CREDS'):
            return
        if random.random() < 0.8 * self.skill:
            new_cred = 'domain_admin' if self.dt_graph.nodes[host].get('value', 0) >= 1000 else 'local_admin'
            state.credentials.add(new_cred)

    def run_simulations(self, num_simulations=100):
        self.attack_graph = nx.DiGraph()
        self.attack_graph.add_node("Attacker", type='Attacker')
        self.stats = self._reset_stats()
        for _ in range(num_simulations):
            self._run_single_intrusion()
        return self.stats, self.attack_graph


# Metrics helper

def extract_metrics(stats):
    total = stats['total'] if stats['total'] > 0 else 1
    sr = stats['succeeded'] / total * 100
    dr = stats['detected'] / total * 100
    fr = stats['failed'] / total * 100
    avg_time, avg_steps = float('nan'), float('nan')
    if stats['succeeded'] > 0:
        all_t = [r['time'] for p in stats['successful_path_details'].values() for r in p]
        all_s = [r['steps'] for p in stats['successful_path_details'].values() for r in p]
        avg_time = np.mean(all_t) if all_t else float('nan')
        avg_steps = np.mean(all_s) if all_s else float('nan')
    return sr, dr, fr, avg_time, avg_steps


# Optuna trial worker

def worker_optuna_trial(args):
    graph_data, params, trial_number, n_reps, sims_per_rep = args
    n_reps = n_reps or 5
    sims_per_rep = sims_per_rep or 50
    G = nx.node_link_graph(graph_data)

    try:
        strat_cls = create_sa_hybrid_strategy(**params)
        sr_list, dr_list, at_list, steps_list = [], [], [], []
        act_list, mp_steps_list = [], []

        for _ in range(n_reps):
            sim = AttackSimulator(G, strat_cls, skill=0.9, detection_rate=0.08)
            stats, _ = sim.run_simulations(num_simulations=sims_per_rep)
            sr, dr, _, avg_time, avg_steps = extract_metrics(stats)
            sr_list.append(sr)
            dr_list.append(dr)
            at_list.append(avg_time)
            steps_list.append(avg_steps)
            strat = sim.strategy
            if hasattr(strat, '_total_activations'):
                act_list.append(strat._total_activations)
                mp_steps_list.append(strat._total_maxprob_steps)

        sr_mean, sr_ci, sr_std = compute_ci(sr_list)
        dr_mean, dr_ci, dr_std = compute_ci(dr_list)
        at_mean, at_ci, at_std = compute_ci(at_list)
        st_mean, st_ci, st_std = compute_ci(steps_list)
        avg_activations = float(np.mean(act_list)) if act_list else 0.0
        avg_mp_steps = float(np.mean(mp_steps_list)) if mp_steps_list else 0.0

        return {
            'trial_number': trial_number,
            'success_rate': sr_mean, 'detection_rate': dr_mean,
            'avg_time': at_mean, 'avg_steps': st_mean,
            'sr_ci': sr_ci, 'dr_ci': dr_ci,
            'at_ci': at_ci, 'steps_ci': st_ci,
            'sr_std': sr_std, 'dr_std': dr_std,
            'maxprob_activations': avg_activations,
            'maxprob_steps': avg_mp_steps,
            'n_reps': n_reps, 'params': params
        }
    except Exception as e:
        return {
            'trial_number': trial_number,
            'success_rate': 0, 'detection_rate': 100,
            'avg_time': float('inf'), 'avg_steps': float('inf'),
            'sr_ci': 0, 'dr_ci': 0,
            'at_ci': float('inf'),
            'steps_ci': float('inf'),
            'sr_std': 0, 'dr_std': 0,
            'maxprob_activations': 0, 'maxprob_steps': 0,
            'n_reps': 0, 'error': str(e), 'params': params
        }


# Parallel simulation worker

def worker_run_simulation(args):
    graph_data, task_spec = args
    G = nx.node_link_graph(graph_data)

    run_idx = task_spec['run_idx']
    n_sims = task_spec['n_sims']

    try:
        if task_spec['type'] == 'baseline':
            strat_name = task_spec['strat_name']
            if 'Stealth' in strat_name:
                strat_cls = StealthStrategy
            else:
                strat_cls = MaxProbabilityStrategy
            sim = AttackSimulator(G, strat_cls, skill=0.9, detection_rate=0.08)
            stats, _ = sim.run_simulations(n_sims)
            sr, dr, fr, avg_time, avg_steps = extract_metrics(stats)

            return {
                'type': 'baseline', 'run_idx': run_idx, 'strat_name': strat_name,
                'config_name': strat_name,
                'success_rate': sr, 'detection_rate': dr, 'failure_rate': fr,
                'avg_time': avg_time, 'avg_steps': avg_steps,
                'avg_activations': 0
            }

        elif task_spec['type'] == 'sahybrid':
            params = task_spec['params']
            config_name = task_spec['config_name']

            strat_cls = create_sa_hybrid_strategy(**params)
            sim = AttackSimulator(G, strat_cls, skill=0.9, detection_rate=0.08)
            stats, _ = sim.run_simulations(n_sims)
            sr, dr, fr, avg_time, avg_steps = extract_metrics(stats)
            avg_act = sim.strategy._total_activations / max(sim.strategy._total_runs, 1)

            return {
                'type': 'sahybrid', 'run_idx': run_idx,
                'config_name': config_name,
                'success_rate': sr, 'detection_rate': dr, 'failure_rate': fr,
                'avg_time': avg_time, 'avg_steps': avg_steps,
                'avg_activations': avg_act, 'params': params
            }

    except Exception as e:
        return {
            'type': task_spec['type'], 'run_idx': run_idx,
            'error': f"{type(e).__name__}: {e}\n{traceback.format_exc()}",
            'strat_name': task_spec.get('strat_name', ''),
            'config_name': task_spec.get('config_name', ''),
        }


# Optuna optimizer

class SAHybridOptimizer:
    def __init__(self, graph_data, n_trials=50, n_reps_per_trial=5, sims_per_rep=50,
                 baseline_maxprob_sr=16.0, baseline_stealth_dr=8.0):
        self.graph_data = graph_data
        self.n_trials = n_trials
        self.n_reps = n_reps_per_trial
        self.sims_per_rep = sims_per_rep
        self.results = []
        self.best_params = None
        self.best_objective = None
        self.study = None
        self.baseline_maxprob_sr = baseline_maxprob_sr
        self.baseline_stealth_dr = baseline_stealth_dr

    def define_objective(self, trial):
        T0 = trial.suggest_float('T0', 0.0, 2.0, step=0.5)
        delta_T = trial.suggest_float('delta_T', 0.5, 3.0, step=0.25)
        T_max = trial.suggest_float('T_max', 5.0, 30.0, step=2.5)
        gamma = trial.suggest_float('gamma', 0.3, 0.9, step=0.1)
        L = trial.suggest_int('L', 1, 5)
        p_func = trial.suggest_categorical('p_func', ['linear', 'sigmoid', 'exponential'])

        params = dict(T0=T0, delta_T=delta_T, T_max=T_max, gamma=gamma, L=L, p_func=p_func)
        result = worker_optuna_trial(
            (self.graph_data, params, trial.number, self.n_reps, self.sims_per_rep)
        )
        self.results.append(result)

        sr = result['success_rate']
        dr = result['detection_rate']
        avg_time = result['avg_time']
        maxprob_act = result.get('maxprob_activations', 0)
        sr_std = result.get('sr_std', 0)
        dr_std = result.get('dr_std', 0)

        if np.isnan(avg_time) or np.isinf(avg_time):
            avg_time = 120.0

        # Priorities: success rate first, then stealth, then speed
        sr_norm = sr / 100.0
        stealth_norm = 1.0 - dr / 100.0
        time_norm = 1.0 / (1.0 + avg_time / 24.0)

        W_SR, W_DR, W_TIME = 10.0, 5.0, 0.5
        objective = W_SR * sr_norm + W_DR * stealth_norm + W_TIME * time_norm

        # Small bonus for beating the baselines
        bonus = 0.0
        if sr > self.baseline_maxprob_sr:
            bonus += 0.3 * (sr - self.baseline_maxprob_sr) / 100.0
        if dr < self.baseline_stealth_dr:
            bonus += 0.2 * (self.baseline_stealth_dr - dr) / 100.0

        # Penalty for extra detection caused by MaxProb deviations
        tradeoff_penalty = 0.0  
        if dr > self.baseline_stealth_dr and maxprob_act > 0:
            excess_dr = (dr - self.baseline_stealth_dr) / 100.0
            tradeoff_penalty = 0.3 * excess_dr * min(maxprob_act / 50.0, 1.0)


        return objective + bonus - tradeoff_penalty

    def optimize(self):
        if not HAS_OPTUNA:
            print("Errore: Optuna non installato. pip install optuna")
            return None

        sampler = TPESampler()
        pruner = MedianPruner(n_startup_trials=5, n_warmup_steps=10)
        self.study = optuna.create_study(
            direction='maximize', sampler=sampler, pruner=pruner,
            study_name='SAHybrid_Optimization'
        )
        self.study.optimize(self.define_objective, n_trials=self.n_trials, show_progress_bar=True)

        best = self.study.best_trial
        self.best_params = best.params
        self.best_objective = best.value
        print(f"Best trial: #{best.number}, obiettivo: {self.best_objective:.4f}")
        return self.study

    def get_top_unique_configs(self, num_top=5):
        sorted_results = sorted(
            [r for r in self.results if 'error' not in r],
            key=lambda x: (
                x['success_rate'],                             # 1st: highest success rate
                -(x['detection_rate']),                        # 2nd: lowest detection rate
                -(max(x.get('avg_time', 0), 0))               # 3rd: shortest time
            ),
            reverse=True
        )
        seen = set()
        unique_top = []
        for r in sorted_results:
            pk = tuple(sorted(r['params'].items()))
            if pk not in seen:
                seen.add(pk)
                unique_top.append(r)
            if len(unique_top) >= num_top:
                break
        return unique_top


# Aggregating results

def aggregate_results(raw_results):
    accum = defaultdict(lambda: defaultdict(list))
    params_map = {}

    for result in raw_results:
        if 'error' in result:
            continue
        config_name = result.get('config_name', result.get('strat_name', 'unknown'))
        for k, v in result.items():
            if k in ('params', 'type', 'run_idx', 'strat_name', 'config_name', 'error'):
                continue
            if isinstance(v, (int, float)):
                accum[config_name][k].append(v)
        if 'params' in result:
            params_map[config_name] = result['params']

    aggregated = {}
    for config, metrics in accum.items():
        agg = {}
        for k, vals in metrics.items():
            arr = np.array(vals, dtype=float)
            arr = arr[~np.isnan(arr)]
            if len(arr) == 0:
                mean, ci, std = np.nan, np.nan, np.nan
            else:
                mean, ci, std = compute_ci(arr)
            agg[k] = {
                'mean': float(mean),
                'ci': float(ci),
                'std': float(std) if not np.isnan(std) else 0.0,
                'n': len(arr)
            }
        if config in params_map:
            agg['params'] = params_map[config]
        aggregated[config] = agg
    return aggregated


# Excel export

def export_to_excel(agg_baseline, agg_sahybrid, comparison_results,
                    n_reps, n_sims, baseline_raw, sahybrid_raw,
                    output_path="risultati_optuna_parallel.xlsx"):
  
    wb = openpyxl.Workbook()

    # ── Styles ──
    header_font = Font(bold=True, size=10, color="FFFFFF")
    fill_blue = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    best_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # light green = best

    thin_border = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )
    center = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    def write_header(ws, headers, row=1):
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = fill_blue
            cell.alignment = center
            cell.border = thin_border

    def style_cell(cell, col_idx, num_fmt=None):
        cell.border = thin_border
        cell.alignment = center if col_idx > 1 else left_align
        if num_fmt:
            cell.number_format = num_fmt

    def auto_width(ws, min_w=8, max_w=30):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max(min_w, min(max_len + 3, max_w))

    # Find the best SAHybrid config
    beating_names = {c['config_name'] for c in comparison_results['beating']}
    best_name = None
    if comparison_results['beating']:
        best_name = comparison_results['beating'][0]['config_name']

    # ── Raw data sheet ──
    config_colors = [
        "D6EAF8",  # light blue
        "D5F5E3",  # light green
        "FCF3CF",  # light yellow
        "FADBD8",  # light pink
        "E8DAEF",  # light purple
        "D6DBDF",  # light grey
        "FDEBD0",  # light orange
        "D4EFDF",  # light mint
        "F9E79F",  # light gold
        "AED6F1",  # sky blue
        "A9DFBF",  # mint green
        "F5CBA7",  # peach
    ]

    ws = wb.active
    ws.title = "Dati Grezzi"

    # Group raw runs by config
    from collections import OrderedDict
    all_raw = baseline_raw + sahybrid_raw
    raw_by_config = OrderedDict()
    for r in all_raw:
        cname = r.get('config_name', r.get('strat_name', 'unknown'))
        if cname not in raw_by_config:
            raw_by_config[cname] = []
        raw_by_config[cname].append(r)

    raw_headers = ["Configurazione", "Run #", "SR (%)", "DR (%)", "Failure (%)",
                   "Tempo (h)", "Steps", "Activations"]
    write_header(ws, raw_headers)

    row_idx = 2
    color_idx = 0
    for cname, runs in raw_by_config.items():
        fill_color = config_colors[color_idx % len(config_colors)]
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        color_idx += 1
        # Sort by run number
        runs_sorted = sorted(runs, key=lambda x: x.get('run_idx', 0))
        for run in runs_sorted:
            vals = [
                cname,
                run.get('run_idx', ''),
                round(run.get('success_rate', 0), 2),
                round(run.get('detection_rate', 0), 2),
                round(run.get('failure_rate', 0), 2),
                round(run.get('avg_time', 0), 4) if not (isinstance(run.get('avg_time'), float) and np.isnan(run.get('avg_time', 0))) else '',
                round(run.get('avg_steps', 0), 1) if not (isinstance(run.get('avg_steps'), float) and np.isnan(run.get('avg_steps', 0))) else '',
                round(run.get('avg_activations', 0), 2),
            ]
            for ci, v in enumerate(vals, 1):
                cell = ws.cell(row=row_idx, column=ci, value=v)
                style_cell(cell, ci, num_fmt='0.00' if isinstance(v, float) else None)
                cell.fill = row_fill
            row_idx += 1
    auto_width(ws)

    # ── Baseline sheet ──
    ws_bl = wb.create_sheet("Baseline")
    bl_headers = ["Strategia", "SR (%)", "SR (±CI)",
                  "DR (%)", "DR (±CI)",
                  "Tempo (h)", "T (±CI)",
                  "Steps", "S (±CI)", "N reps"]
    write_header(ws_bl, bl_headers)
    for row_idx, (name, m) in enumerate(agg_baseline.items(), 2):
        sr = m.get('success_rate', {})
        dr = m.get('detection_rate', {})
        at = m.get('avg_time', {})
        st = m.get('avg_steps', {})
        vals = [
            name,
            sr.get('mean', 0), sr.get('ci', 0),
            dr.get('mean', 0), dr.get('ci', 0),
            at.get('mean', 0), at.get('ci', 0),
            st.get('mean', 0), st.get('ci', 0),
            sr.get('n', 0)
        ]
        for ci, v in enumerate(vals, 1):
            cell = ws_bl.cell(row=row_idx, column=ci, value=v)
            style_cell(cell, ci, num_fmt='0.00' if isinstance(v, float) else None)
    auto_width(ws_bl)

    # ── SAHybrid configs sheet ──
    ws_sa = wb.create_sheet("SAHybrid Configs")
    sa_headers = ["Config", "T0", "delta_T", "T_max", "gamma", "L", "p(T)",
                  "SR (%)", "SR (±CI)",
                  "DR (%)", "DR (±CI)",
                  "Tempo (h)", "T (±CI)",
                  "Steps", "S (±CI)",
                  "Acts/run", "N reps", "Migliore?"]
    write_header(ws_sa, sa_headers)
    row_idx = 2
    for name, m in agg_sahybrid.items():
        sr = m.get('success_rate', {})
        dr = m.get('detection_rate', {})
        at = m.get('avg_time', {})
        st = m.get('avg_steps', {})
        acts = m.get('avg_activations', {})
        params = m.get('params', {})
        beats = name in beating_names
        vals = [
            name,
            params.get('T0', ''), params.get('delta_T', ''), params.get('T_max', ''),
            params.get('gamma', ''), params.get('L', ''), params.get('p_func', ''),
            sr.get('mean', 0), sr.get('ci', 0),
            dr.get('mean', 0), dr.get('ci', 0),
            at.get('mean', 0), at.get('ci', 0),
            st.get('mean', 0), st.get('ci', 0),
            acts.get('mean', 0) if isinstance(acts, dict) else acts,
            sr.get('n', 0),
            "SI" if beats else "NO"
        ]
        for ci, v in enumerate(vals, 1):
            cell = ws_sa.cell(row=row_idx, column=ci, value=v)
            style_cell(cell, ci, num_fmt='0.00' if isinstance(v, float) else None)
            # Highlight only the best row
            if name == best_name:
                cell.fill = best_fill
        row_idx += 1
    auto_width(ws_sa, min_w=7, max_w=25)

    # ── Comparison vs MaxProb sheet ──
    ws_cmp = wb.create_sheet("Confronto vs MaxProb")
    cmp_headers = ["Config", "T0", "delta_T", "T_max", "gamma", "L", "p(T)",
                   "SR Mean", "SR (±CI)",
                   "DR Mean", "DR (±CI)",
                   "MaxP SR Mean", "Delta SR (%)", "Risultato"]
    write_header(ws_cmp, cmp_headers)

    # MaxProb reference values
    mp_agg = agg_baseline.get('MaxProbability', {})
    mp_sr = mp_agg.get('success_rate', {})
    mp_dr = mp_agg.get('detection_rate', {})
    mp_sr_mean = mp_sr.get('mean', 0)

    # List every SAHybrid config, best SR first
    sahybrid_sorted = sorted(
        agg_sahybrid.items(),
        key=lambda x: x[1].get('success_rate', {}).get('mean', 0),
        reverse=True
    )

    better_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    inferior_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    row_idx = 2
    for name, m in sahybrid_sorted:
        sr = m.get('success_rate', {})
        dr = m.get('detection_rate', {})
        params = m.get('params', {})
        sr_m = sr.get('mean', 0)
        dr_m = dr.get('mean', 0)
        delta = sr_m - mp_sr_mean

        # Does it beat MaxProb?
        if sr_m > mp_sr_mean:
            tipo = "SUPERIORE"
        elif abs(sr_m - mp_sr_mean) < 0.01:
            tipo = "PARI"
        else:
            tipo = "INFERIORE"

        vals = [
            name,
            params.get('T0', ''), params.get('delta_T', ''), params.get('T_max', ''),
            params.get('gamma', ''), params.get('L', ''), params.get('p_func', ''),
            sr_m, sr.get('ci', 0),
            dr_m, dr.get('ci', 0),
            mp_sr_mean, delta, tipo
        ]

        # Color the row by result
        if tipo == "SUPERIORE":
            row_fill = better_fill
        elif tipo == "INFERIORE":
            row_fill = inferior_fill
        else:
            row_fill = None

        for ci, v in enumerate(vals, 1):
            cell = ws_cmp.cell(row=row_idx, column=ci, value=v)
            style_cell(cell, ci, num_fmt='0.00' if isinstance(v, float) else None)
            if row_fill:
                cell.fill = row_fill
        row_idx += 1

    # Reference row: MaxProbability
    row_idx += 1
    ref_vals = [
        "MaxProbability (REF)", "", "", "", "", "", "",
        mp_sr_mean, mp_sr.get('ci', 0),
        mp_dr.get('mean', 0), mp_dr.get('ci', 0),
        "", "", "BASELINE"
    ]
    ref_fill = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid")
    for ci, v in enumerate(ref_vals, 1):
        cell = ws_cmp.cell(row=row_idx, column=ci, value=v)
        cell.border = thin_border
        cell.alignment = center
        cell.font = Font(bold=True)
        cell.fill = ref_fill

    auto_width(ws_cmp, min_w=7, max_w=25)

    wb.save(output_path)
    print(f"Salvato: {output_path}")
    return wb


# Main

def main():
    TRIALS      = 300      # Optuna trials
    SIMS        = 30000    # Simulations per repetition in the final evaluation
    REPS        = 100      # Repetitions used for the confidence intervals
    OPTUNA_REPS = 20           # Repetitions per Optuna trial (kept low for speed)
    OPTUNA_SIMS = 500      # Simulations per repetition during tuning
    TOP_EVAL    = 5       # Top Optuna configs to re-evaluate
    SKIP_PFUNC  = False   # True = skip the p(T) shape comparison
    MAXWORKERS = None    # Max parallel workers (None = CPU count)

    start_time = datetime.now()
    timestamp = start_time.strftime("%Y%m%d_%H%M%S")
    print(f"Optuna pipeline | {TRIALS} trials, {REPS} reps x {SIMS} sim")

    nest_asyncio.apply()
    
    # Config
    NVD_API_KEY = None 
    CSV_FILE_PATH = 'glpi.csv'
    
    # Build the digital twin
    dt = DigitalTwin()
    try:
        dt.load_from_csv(CSV_FILE_PATH)
    except FileNotFoundError:
        print(f"ERRORE: il file '{CSV_FILE_PATH}' non è stato trovato.")
        exit(1)
    
    # Add CVE data
    enricher = CVEEnricher(dt.get_graph(), api_key=NVD_API_KEY)
    asyncio.run(enricher.run_enrichment())
    graph = dt.get_graph()
    graph_data = nx.node_link_data(graph)

    # PHASE 2: Optuna tuning
    quick_reps = max(3, OPTUNA_REPS)
    quick_sims = OPTUNA_SIMS

    # Quick MaxProb baseline
    quick_mp_sr_list = []
    for _ in range(quick_reps):
        sim = AttackSimulator(graph, MaxProbabilityStrategy, skill=0.9, detection_rate=0.08)
        stats, _ = sim.run_simulations(quick_sims)
        quick_mp_sr_list.append(stats['succeeded'] / stats['total'] * 100 if stats['total'] > 0 else 0)
    quick_mp_sr = float(np.mean(quick_mp_sr_list))

    # Quick Stealth baseline
    quick_st_dr_list = []
    for _ in range(quick_reps):
        sim = AttackSimulator(graph, StealthStrategy, skill=0.9, detection_rate=0.08)
        stats, _ = sim.run_simulations(quick_sims)
        quick_st_dr_list.append(stats['detected'] / stats['total'] * 100 if stats['total'] > 0 else 0)
    quick_st_dr = float(np.mean(quick_st_dr_list))

    optimizer = SAHybridOptimizer(
        graph_data, n_trials=TRIALS,
        n_reps_per_trial=OPTUNA_REPS, sims_per_rep=OPTUNA_SIMS,
        baseline_maxprob_sr=quick_mp_sr,
        baseline_stealth_dr=quick_st_dr
    )
    optimizer.optimize()

    if not optimizer.best_params:
        print("Errore: ottimizzazione fallita, nessun risultato.")
        return

    top_configs = optimizer.get_top_unique_configs(num_top=TOP_EVAL)
    print(f"Top {len(top_configs)} config selezionate per valutazione robusta.")

    # Top configs, plus p_func variants of the best one
    configs_to_evaluate = []
    for idx, cfg in enumerate(top_configs, 1):
        label = f"Top{idx}_T{cfg['trial_number']}"
        configs_to_evaluate.append((label, cfg['params']))

    if not SKIP_PFUNC:
        best_p = optimizer.best_params
        for pf in ['linear', 'sigmoid', 'exponential']:
            variant_params = {**best_p, 'p_func': pf}
            variant_label = f"BestVariant_{pf}"
            # Skip duplicates
            already = any(
                json.dumps(p, sort_keys=True) == json.dumps(variant_params, sort_keys=True)
                for _, p in configs_to_evaluate
            )
            if not already:
                configs_to_evaluate.append((variant_label, variant_params))

    # PHASE 3: Parallel evaluation
    tasks = []

    # Baselines
    for strat_name in ['MaxProbability', 'Stealth']:
        for rep_idx in range(1, REPS + 1):
            tasks.append((graph_data, {
                'type': 'baseline', 'run_idx': rep_idx,
                'strat_name': strat_name, 'n_sims': SIMS
            }))

    # SAHybrid configs
    for label, params in configs_to_evaluate:
        for rep_idx in range(1, REPS + 1):
            tasks.append((graph_data, {
                'type': 'sahybrid', 'run_idx': rep_idx,
                'config_name': label, 'params': params, 'n_sims': SIMS
            }))

    total_tasks = len(tasks)
    nworkers = MAXWORKERS or min(multiprocessing.cpu_count(), total_tasks)
    print(f"Valutazione parallela: {total_tasks} task, {nworkers} worker")

    all_raw_results = []
    completed = 0

    with ProcessPoolExecutor(max_workers=nworkers) as executor:
        futures = {executor.submit(worker_run_simulation, task): task for task in tasks}
        for future in as_completed(futures):
            all_raw_results.append(future.result())
            completed += 1
            if completed % 50 == 0 or completed == total_tasks:
                print(f"  [{completed}/{total_tasks}] completate")

    # PHASE 4: Aggregate with CIs
    baseline_raw = [r for r in all_raw_results if r.get('type') == 'baseline' and 'error' not in r]
    sahybrid_raw = [r for r in all_raw_results if r.get('type') == 'sahybrid' and 'error' not in r]

    agg_baseline = aggregate_results(baseline_raw)
    agg_sahybrid = aggregate_results(sahybrid_raw)

    # PHASE 5: Compare against MaxProb
    mp_agg = agg_baseline.get('MaxProbability', {})
    mp_sr = mp_agg.get('success_rate', {})
    mp_dr = mp_agg.get('detection_rate', {})
    mp_sr_mean = mp_sr.get('mean', 0)

    beating = []

    for name, m in agg_sahybrid.items():
        sr = m.get('success_rate', {})
        dr = m.get('detection_rate', {})
        sr_mean = sr.get('mean', 0)
        dr_mean = dr.get('mean', 0)

        entry = {
            'config_name': name,
            'params': m.get('params', {}),
            'sr_mean': sr_mean,
            'dr_mean': dr_mean,
            'delta_sr': sr_mean - mp_sr_mean,
            'delta_dr': dr_mean - mp_dr.get('mean', 0),
        }

        # Mean SR only, no significance test
        if sr_mean > mp_sr_mean:
            beating.append(entry)

    # Sort by SR (high), then DR (low)
    beating.sort(key=lambda x: (x['sr_mean'], -x['dr_mean']), reverse=True)

    comparison_results = {
        'beating': beating,
    }

    # PHASE 6: Save
    print(f"Confronto: {len(beating)} configurazioni con media SR superiore a MaxProb.")
    excel_output = f"risultati_optuna_parallel_{timestamp}.xlsx"
    export_to_excel(
        agg_baseline, agg_sahybrid, comparison_results,
        REPS, SIMS, baseline_raw, sahybrid_raw,
        output_path=excel_output
    )

    elapsed = datetime.now() - start_time
    print(f"Completato in {elapsed.total_seconds()/60:.1f} min.")


if __name__ == '__main__':
    main()