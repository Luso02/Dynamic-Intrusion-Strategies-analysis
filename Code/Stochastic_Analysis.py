

import csv
import networkx as nx
import re
import ipaddress
from collections import defaultdict, Counter
import json
import os
import asyncio
import aiohttp
import nest_asyncio
import random
import math
import numpy as np
from scipy.optimize import curve_fit
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime
import multiprocessing
import traceback

# For Excel
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    print("⚠️  openpyxl non installato: pip install openpyxl  — l'export Excel verrà saltato.")


# Core classes

class Asset:
    def __init__(self, name, **kwargs):
        self.name = name
        self.attributes = kwargs
    def __repr__(self):
        return f"{self.__class__.__name__}(name='{self.name}')"

class Host(Asset):
    pass

class Software(Asset):
    def __init__(self, name, **kwargs):
        super().__init__(name, **kwargs)
        self.base_name = kwargs.get('base_name')
        self.version = kwargs.get('version')

class VirtualMachine(Host):
    pass


# Pulls CVE data from the NVD API and attaches scores to software nodes
class CVEEnricher:
    def __init__(self, graph, api_key=None, cache_file='cve_cache.json'):
        self.graph = graph
        self.api_key = api_key
        self.cache_file = cache_file
        self.cve_cache = self._load_cache()
        print(f"🔎 CVEEnricher inizializzato. {len(self.cve_cache)} voci in cache.")

    def _load_cache(self):
        if os.path.exists(self.cache_file):
            with open(self.cache_file, 'r') as f:
                return json.load(f)
        return {}

    def _save_cache(self):
        with open(self.cache_file, 'w') as f:
            json.dump(self.cve_cache, f, indent=2)

    def _normalize_sw_name(self, name):
        # Keep only the part before the first dash for the NVD search
        return name.split('-')[0]

    async def _fetch_cve(self, session, sw_name, clean_version, original_key):
        if original_key in self.cve_cache:
            return original_key, self.cve_cache[original_key], '.'
        search_name = self._normalize_sw_name(sw_name)
        headers = {'apiKey': self.api_key} if self.api_key else {}
        query = f"keywordSearch={search_name} {clean_version}"
        url = f"https://services.nvd.nist.gov/rest/json/cves/2.0?{query}"
        try:
            async with session.get(url, headers=headers, timeout=40) as response:
                response.raise_for_status()
                data = await response.json()
                vulnerabilities = []
                if 'vulnerabilities' in data:
                    for item in data['vulnerabilities']:
                        cve = item['cve']
                        # -1 means no CVSS score available
                        score = -1.0
                        # Prefer CVSS v3.1, fall back to v2
                        if 'cvssMetricV31' in cve['metrics']:
                            score = cve['metrics']['cvssMetricV31'][0]['cvssData']['baseScore']
                        elif 'cvssMetricV2' in cve['metrics']:
                            score = cve['metrics']['cvssMetricV2'][0]['cvssData']['baseScore']
                        vulnerabilities.append({'id': cve['id'], 'score': score})
                self.cve_cache[original_key] = vulnerabilities
                feedback_char = 'V' if vulnerabilities else '.'
                return original_key, vulnerabilities, feedback_char
        except Exception:
            # On failure, cache an empty result so we don't retry
            self.cve_cache[original_key] = []
            return original_key, [], 'E'

    async def run_enrichment(self):
        print("\n📡 Inizio arricchimento dati CVE...")
        unique_software = {
            f"{data.get('base_name', '').lower()}|{data.get('version', '')}": (
                data.get('base_name'),
                re.match(r'[\d\.:]+', data.get('version', '')).group(0)
                if re.match(r'[\d\.:]+', data.get('version', ''))
                else data.get('version', '')
            )
            for _, data in self.graph.nodes(data=True)
            if data.get('type') == 'Software' and data.get('base_name') and data.get('version')
        }
        print(f"   -> {len(unique_software)} software unici.")
        items_to_fetch = [
            (key, sw_name, clean_version)
            for key, (sw_name, clean_version) in unique_software.items()
            if key not in self.cve_cache
        ]
        if not items_to_fetch:
            print("   -> Tutto in cache.")
        else:
            print(f"   -> {len(items_to_fetch)} da scaricare...")
            async with aiohttp.ClientSession() as session:
                results = await asyncio.gather(
                    *[self._fetch_cve(session, sw, ver, key) for key, sw, ver in items_to_fetch]
                )
                print(''.join(fb for _, _, fb in results), flush=True)
        for node_name, data in self.graph.nodes(data=True):
            if data.get('type') == 'Software':
                key = f"{data.get('base_name', '').lower()}|{data.get('version', '')}"
                if (vulns := self.cve_cache.get(key)):
                    self.graph.nodes[node_name]['vulnerabilities'] = vulns
                    self.graph.nodes[node_name]['max_cvss'] = max(
                        [v['score'] for v in vulns if v['score'] >= 0] or [0.0]
                    )
        self._save_cache()
        print("✅ Arricchimento CVE completato.")


# Graph of hosts, VMs and software built from a GLPI CSV export
class DigitalTwin:
    def __init__(self):
        self.graph = nx.DiGraph()
        self.assets = {}
        self.discovered_subnets = []

    def _parse_multiline_data(self, data: str):
        # GLPI packs several values in one cell, separated by <br>
        return [item.strip() for item in data.split('<br>') if item.strip()] if isinstance(data, str) else []

    def _parse_software_data(self, data: str):
        return [
            {'name': parts[0].strip(), 'version': parts[1].strip()}
            for line in self._parse_multiline_data(data)
            if len(parts := line.rsplit(' - ', 1)) == 2
        ]

    def _discover_subnets_from_rows(self, all_rows):
        print("🔎 Scoperta automatica delle sottoreti...")
        subnet_prefixes = set()
        for row in all_rows:
            for ip_str in self._parse_multiline_data(row.get('Networking - IP', '')):
                try:
                    ip_addr = ipaddress.ip_address(ip_str)
                    # Keep private IPv4 addresses only, grouped into /24 subnets
                    if not (ip_addr.is_loopback or ip_addr.version == 6 or not ip_addr.is_private):
                        subnet_prefixes.add('.'.join(ip_str.split('.')[:3]))
                except ValueError:
                    continue
        self.discovered_subnets = [
            ipaddress.ip_network(f"{prefix}.0/24") for prefix in sorted(subnet_prefixes)
        ]
        print(f"✅ Sottoreti scoperte: {len(self.discovered_subnets)}")

    def _get_subnet_for_ips(self, ips):
        for ip_str in ips:
            try:
                subnet = next(
                    (str(s) for s in self.discovered_subnets if ipaddress.ip_address(ip_str) in s), None
                )
                if subnet:
                    return subnet
            except ValueError:
                continue
        # No known subnet: treated as external (attacker entry points)
        return 'Unknown/External'

    def _add_or_get_asset(self, name, asset_class, **kwargs):
        if name not in self.assets:
            self.assets[name] = asset_class(name, **kwargs)
            self.graph.add_node(name, type=asset_class.__name__, **kwargs)
        # Asset already exists: just update its attributes
        elif kwargs:
            nx.set_node_attributes(self.graph, {name: kwargs})
        return self.assets[name]

    def load_from_csv(self, file_path, delimiter=';'):
        print(f"📄 Lettura: {file_path}")
        with open(file_path, mode='r', encoding='utf-8-sig') as infile:
            next(infile)
            all_rows = list(csv.DictReader(infile, delimiter=delimiter))
        self._discover_subnets_from_rows(all_rows)
        print("🏗️  Costruzione del grafo...")
        for row in all_rows:
            if not (host_name := row.get('Name')):
                continue
            ips = self._parse_multiline_data(row.get('Networking - IP', ''))
            subnet = self._get_subnet_for_ips(ips)
            host_attributes = {
                'os': row.get('Operating System - Name'),
                'asset_type': row.get('Type'), 'ips': ips, 'subnet': subnet
            }
            self._add_or_get_asset(
                host_name, VirtualMachine if host_attributes['asset_type'] == 'VM' else Host,
                **host_attributes
            )
            for vm_name in self._parse_multiline_data(row.get('Virtual machines - Name', '')):
                self._add_or_get_asset(vm_name, VirtualMachine, os=host_attributes['os'], subnet=subnet)
                self.graph.add_edge(host_name, vm_name, relationship='HOSTS')
            for sw in self._parse_software_data(row.get('Software - Name', '')):
                sw_unique_name = f"{sw['name']} ({sw['version']})"
                self._add_or_get_asset(sw_unique_name, Software, base_name=sw['name'], version=sw['version'])
                self.graph.add_edge(host_name, sw_unique_name, relationship='INSTALLS')
        print("✅ Grafo costruito.")

    def get_graph(self):
        return self.graph


# Attacker state and strategies

class AttackerState:
    def __init__(self, initial_host):
        # What the attacker owns so far, and how noisy the intrusion has been
        self.compromised_hosts = {initial_host: {'priv_level': 'user'}}
        self.credentials = set()
        self.current_time_hours = 0.0
        self.path = ['Attacker', initial_host]
        self.detected = False
        self.goal_achieved = False
        self.noise_level = 0.0


class AttackStrategy:
    def __init__(self, simulator):
        self.sim = simulator
    def choose_next_action(self, state, current_host):
        raise NotImplementedError
    def choose_target(self, state, current_host, available_targets):
        raise NotImplementedError
    def __str__(self):
        return self.__class__.__name__


class StealthStrategy(AttackStrategy):
    def __init__(self, simulator, noise_threshold=10.0):
        super().__init__(simulator)
        self.noise_threshold = noise_threshold
    def choose_next_action(self, state, host):
        # Too noisy: lay low and let the noise decay
        if state.noise_level > self.noise_threshold:
            return 'WAIT'
        priv_level = state.compromised_hosts[host]['priv_level']
        if 'domain_admin' in state.credentials or 'local_admin' in state.credentials:
            return 'LATERAL_MOVE'
        if priv_level == 'user':
            return 'PRIV_ESC'
        if priv_level == 'admin':
            return 'DUMP_CREDS'
        return 'LATERAL_MOVE'
    def choose_target(self, state, current_host, targets):
        if not targets:
            return None
        # Least valuable host first, to draw less attention
        return min(targets, key=lambda t: self.sim.dt_graph.nodes[t].get('value', 1))


class MaxProbabilityStrategy(AttackStrategy):
    def choose_next_action(self, state, host):
        if state.compromised_hosts[host]['priv_level'] == 'user':
            return 'PRIV_ESC'
        # On a domain controller, go for domain admin credentials
        if 'domain_admin' not in state.credentials and self.sim.dt_graph.nodes[host].get('value', 0) >= 1000:
            return 'DUMP_CREDS'
        if 'local_admin' not in state.credentials:
            return 'DUMP_CREDS'
        return 'LATERAL_MOVE'
    def _get_max_cvss_for_host(self, target_host):
        vulns = [
            self.sim.dt_graph.nodes[s].get('max_cvss', 0.0)
            for s in self.sim.dt_graph.successors(target_host)
            if self.sim.dt_graph.nodes[s].get('type') == 'Software'
            and self.sim.dt_graph.nodes[s].get('vulnerabilities')
        ]
        return max(vulns) if vulns else 0.0
    def choose_target(self, state, current_host, targets):
        if not targets:
            return None
        # Go for the host with the most vulnerable software
        return max(targets, key=lambda t: self._get_max_cvss_for_host(t))


class StochasticPerturbationStrategy(AttackStrategy):
    """
    Stochastic hybrid: MaxProbability is the base, but at each step it
    deviates with probability p, making the attack less predictable and
    lowering the detection rate.
    """
    def __init__(self, simulator, perturbation_prob=0.01, deviation_mode='stealth'):
        super().__init__(simulator)
        self.perturbation_prob = perturbation_prob  # p: e.g. 0.01, 0.03, 0.05
        self.deviation_mode = deviation_mode  # 'stealth', 'random', 'mixed'
        self.maxprob_strategy = MaxProbabilityStrategy(simulator)
        self.stealth_strategy = StealthStrategy(simulator, noise_threshold=10.0)
        self._deviations_count = 0
        self._total_actions = 0
        
    def choose_next_action(self, state, host):
        self._total_actions += 1
        # With probability p, deviate from MaxProb
        if random.random() < self.perturbation_prob:
            self._deviations_count += 1
            if self.deviation_mode == 'stealth':
                # Take a Stealth action
                return self.stealth_strategy.choose_next_action(state, host)
            elif self.deviation_mode == 'random':
                # Any random action
                return random.choice(['PRIV_ESC', 'DUMP_CREDS', 'LATERAL_MOVE', 'WAIT'])
            elif self.deviation_mode == 'mixed':
                # Half Stealth, half random
                if random.random() < 0.5:
                    return self.stealth_strategy.choose_next_action(state, host)
                else:
                    return random.choice(['PRIV_ESC', 'DUMP_CREDS', 'LATERAL_MOVE', 'WAIT'])
        
        # Otherwise follow MaxProb
        return self.maxprob_strategy.choose_next_action(state, host)
    
    def choose_target(self, state, current_host, targets):
        # Targets always come from MaxProb
        return self.maxprob_strategy.choose_target(state, current_host, targets)
    
    def get_deviation_rate(self):
        """Share of actions (%) that deviated."""
        if self._total_actions == 0:
            return 0.0
        return (self._deviations_count / self._total_actions) * 100
    
    def __str__(self):
        p_percent = int(self.perturbation_prob * 100)
        return f"StochasticPerturbation_p{p_percent}_{self.deviation_mode}"

# Returns a strategy class with p and mode fixed (the simulator takes a class)
def create_stochastic_variant(p, mode='stealth'):
    class VariantStrategy(StochasticPerturbationStrategy):
        def __init__(self, sim):
            super().__init__(sim, perturbation_prob=p, deviation_mode=mode)
    return VariantStrategy
 

# Simulation engine

class AttackSimulator:
    # Time range (hours) and detection noise for each action
    TIME_MODEL = {
        'SCAN': (1.0, 1.0), 'EXPLOIT': (0.1, 2), 'PRIV_ESC': (0.2, 3),
        'DUMP_CREDS': (0.1, 0.5), 'LOGIN': (0.01, 0.1)
    }
    NOISE_MODEL = {'SCAN': 1.5, 'EXPLOIT': 2.0, 'PRIV_ESC': 1.0, 'DUMP_CREDS': 1.2, 'LOGIN': 0.2}

    def __init__(self, dt_graph, strategy_class, skill=0.8, detection_rate=0.05):
        self.dt_graph = dt_graph
        self.strategy = strategy_class(self)
        self.skill = skill
        self.detection_rate = detection_rate
        self.attack_graph = nx.DiGraph()
        self._assign_asset_values()
        self.stats = self._reset_stats()

    def _reset_stats(self):
        return {
            'total': 0, 'succeeded': 0, 'failed': 0, 'detected': 0,
            'actions': defaultdict(int), 'exploited_cves': defaultdict(int),
            'compromised_nodes': defaultdict(int),
            'successful_path_details': defaultdict(list),
            'detected_paths': []
        }

    def _assign_asset_values(self):
        # Host value from name keywords; domain controllers (1000) are the goal
        value_map = {'dc': 1000, 'db': 500, 'sql': 500, 'fileserver': 200, 'admin': 100}
        for node, data in self.dt_graph.nodes(data=True):
            if data.get('type') in ['Host', 'VirtualMachine']:
                self.dt_graph.nodes[node]['value'] = max(
                    [1] + [v for keyword, v in value_map.items() if keyword in node.lower()]
                )

    def _select_entry_point(self):
        # Entry point: external hosts, weighted by value, boosted if they run a web server
        web_keywords = ['apache', 'nginx', 'iis', 'httpd', 'tomcat']
        candidates = [
            (host, 1 + data.get('value', 1) + 10 * any(
                any(key in sw_name for key in web_keywords)
                for sw in self.dt_graph.successors(host)
                if (sw_name := self.dt_graph.nodes[sw].get('base_name', '').lower())
            ))
            for host, data in self.dt_graph.nodes(data=True)
            if data.get('type') in ['Host', 'VirtualMachine'] and data.get('subnet') == 'Unknown/External'
        ]
        if not candidates:
            return random.choice(
                [n for n, d in self.dt_graph.nodes(data=True) if d.get('type') in ['Host', 'VirtualMachine']]
            )
        return random.choices(*zip(*candidates), k=1)[0]

    def _get_noise_for_action(self, action):
        base_noise = self.NOISE_MODEL.get(action, 1.0)
        # Strategies can optionally scale their noise
        multiplier = getattr(self.strategy, 'noise_multiplier', 1.0)
        return base_noise * multiplier

    def _run_single_intrusion(self):
        self.stats['total'] += 1
        entry_host = self._select_entry_point()
        sw, vuln = self._attempt_exploit_cve(entry_host)
        if not sw:
            self.stats['failed'] += 1
            return
        state = AttackerState(entry_host)
        state.current_time_hours += self._get_time('EXPLOIT')
        state.noise_level += self._get_noise_for_action('EXPLOIT')

        # Keep going until caught, done, or out of time (120 h)
        while not state.detected and not state.goal_achieved and state.current_time_hours < 120:
            host = state.path[-1]
            # Reached a domain controller: goal achieved
            if self.dt_graph.nodes[host].get('value', 0) >= 1000:
                state.goal_achieved = True
                continue
            action = self.strategy.choose_next_action(state, host)
            if action == 'PRIV_ESC':
                self._attempt_privesc(state, host)
            elif action == 'DUMP_CREDS':
                self._attempt_credential_dump(state, host)
            elif action == 'LATERAL_MOVE':
                self._attempt_lateral_move(state, host)
            # Lay low: time passes and the noise fades
            elif action == 'WAIT':
                state.current_time_hours += 12
                state.noise_level *= 0.75

        # Success = reached the goal or got past the entry host
        if state.goal_achieved or len(state.compromised_hosts) > 1:
            self.stats['succeeded'] += 1
            path_tuple = tuple(state.path)
            self.stats['successful_path_details'][path_tuple].append({
                'time': state.current_time_hours, 'steps': len(state.path)
            })
        elif state.detected:
            self.stats['detected'] += 1
            self.stats['detected_paths'].append(tuple(state.path))
        else:
            self.stats['failed'] += 1

    def _attempt_lateral_move(self, state, current_host):
        state.current_time_hours += self._get_time('SCAN')
        state.noise_level += self._get_noise_for_action('SCAN')
        if self._check_detection(state, 'SCAN'):
            return
        subnet = self.dt_graph.nodes[current_host].get('subnet')
        targets = [
            n for n, d in self.dt_graph.nodes(data=True)
            if d.get('type') in ['Host', 'VirtualMachine']
            and d.get('subnet') == subnet and n not in state.compromised_hosts
        ]
        target_host = self.strategy.choose_target(state, current_host, targets)
        if not target_host:
            return
        moved, sw, vuln = False, None, None
        # With admin credentials, log in instead of exploiting
        if 'domain_admin' in state.credentials and random.random() < 0.95:
            moved = True
        elif 'local_admin' in state.credentials and random.random() < 0.6:
            moved = True
        if moved:
            state.current_time_hours += self._get_time('LOGIN')
            state.noise_level += self._get_noise_for_action('LOGIN')
        else:
            state.current_time_hours += self._get_time('EXPLOIT')
            state.noise_level += self._get_noise_for_action('EXPLOIT')
            if self._check_detection(state, 'EXPLOIT'):
                return
            sw, vuln = self._attempt_exploit_cve(target_host)
            if sw:
                moved = True
        if moved:
            state.compromised_hosts[target_host] = {'priv_level': 'user'}
            state.path.append(target_host)

    def _get_time(self, action):
        # Higher skill means faster actions
        return random.uniform(*self.TIME_MODEL[action]) / self.skill

    def _check_detection(self, state, action):
        # Noisier actions are more likely to be detected
        if random.random() < self.detection_rate * self._get_noise_for_action(action):
            state.detected = True
            return True
        return False

    def _attempt_exploit_cve(self, target_host):
        vulns = [
            (s, max(self.dt_graph.nodes[s].get('vulnerabilities', []), key=lambda v: v['score']))
            for s in self.dt_graph.successors(target_host)
            if self.dt_graph.nodes[s].get('type') == 'Software'
            and self.dt_graph.nodes[s].get('vulnerabilities')
        ]
        if not vulns:
            return None, None
        sw_node, best_vuln = max(vulns, key=lambda item: item[1]['score'])
        # Success chance scales with the CVSS score and the attacker's skill
        if random.random() < (best_vuln['score'] / 10.0) * self.skill:
            return sw_node, best_vuln
        return None, None

    def _attempt_privesc(self, state, host):
        state.current_time_hours += self._get_time('PRIV_ESC')
        state.noise_level += self._get_noise_for_action('PRIV_ESC')
        if self._check_detection(state, 'PRIV_ESC'):
            return
        if random.random() < 0.7 * self.skill:
            state.compromised_hosts[host]['priv_level'] = 'admin'

    def _attempt_credential_dump(self, state, host):
        if state.compromised_hosts[host]['priv_level'] != 'admin':
            return
        state.current_time_hours += self._get_time('DUMP_CREDS')
        state.noise_level += self._get_noise_for_action('DUMP_CREDS')
        if self._check_detection(state, 'DUMP_CREDS'):
            return
        if random.random() < 0.8 * self.skill:
            new_cred = 'domain_admin' if self.dt_graph.nodes[host].get('value', 0) >= 1000 else 'local_admin'
            state.credentials.add(new_cred)

    def run_simulations(self, num_simulations=100):
        self.attack_graph = nx.DiGraph()
        self.attack_graph.add_node("Attacker", type='Attacker')
        self.stats = self._reset_stats()
        for _ in range(num_simulations):
            self._run_single_intrusion()
        return self.stats, self.attack_graph


# Parallel worker

def _worker_run_simulation(args):
    """
    Runs in a separate process.
    Takes (graph_data, task_spec):
      graph_data: serialized NetworkX graph (nodes and edges)
      task_spec:  task type ('baseline' or 'sensitivity'), parameters, etc.

    Returns a dict with the results.
    """
    graph_data, task_spec = args

    # Rebuild the graph in the child process
    G = nx.DiGraph()
    for name, attrs in graph_data['nodes']:
        # Deserialize vulnerabilities if needed
        for key in ['vulnerabilities', 'ips']:
            if key in attrs and isinstance(attrs[key], str):
                attrs[key] = json.loads(attrs[key])
        G.add_node(name, **attrs)
    for u, v, attrs in graph_data['edges']:
        G.add_edge(u, v, **attrs)

    run_idx = task_spec['run_idx']
    n_sims = task_spec['n_sims']

    try:
        if task_spec['type'] == 'baseline':
            strat_name = task_spec['strat_name']
            strat_cls = StealthStrategy if 'Stealth' in strat_name else MaxProbabilityStrategy
            sim = AttackSimulator(G, strat_cls, skill=0.9, detection_rate=0.08)
            stats, _ = sim.run_simulations(n_sims)

            sr = stats['succeeded'] / stats['total'] * 100
            dr = stats['detected'] / stats['total'] * 100
            fr = stats['failed'] / stats['total'] * 100
            avg_time, ci_time, avg_steps, ci_steps = float('nan'), 0, float('nan'), 0
            if stats['succeeded'] > 0:
                all_t = [r['time'] for p in stats['successful_path_details'].values() for r in p]
                all_s = [r['steps'] for p in stats['successful_path_details'].values() for r in p]
                avg_time = np.mean(all_t)
                ci_time = (1.96 * np.std(all_t, ddof=1) / np.sqrt(len(all_t))) if len(all_t) > 1 else 0
                avg_steps = np.mean(all_s)
                ci_steps = (1.96 * np.std(all_s, ddof=1) / np.sqrt(len(all_s))) if len(all_s) > 1 else 0

            return {
                'type': 'baseline', 'run_idx': run_idx, 'strat_name': strat_name,
                'success_rate': sr, 'detection_rate': dr, 'failure_rate': fr,
                'avg_time': avg_time, 'ci_time': ci_time,
                'avg_steps': avg_steps, 'ci_steps': ci_steps,
                'avg_activations': 0
            }

        elif task_spec['type'] == 'sensitivity':
            config_name = task_spec['config_name']
            param_key = task_spec['param_key']
            
            # StochasticPerturbation support
            if 'stochastic_p' in task_spec:
                # Stochastic perturbation task
                p = task_spec['stochastic_p']
                deviation_mode = task_spec.get('deviation_mode', 'stealth')
                
                # Subclass with p and mode baked in
                class StocasticVariant(StochasticPerturbationStrategy):
                    def __init__(self, sim):
                        super().__init__(sim, perturbation_prob=p, deviation_mode=deviation_mode)
                
                strat_cls = StocasticVariant
                avg_activations_label = 'deviation_rate'  # This strategy reports the deviation rate
           
            
            sim = AttackSimulator(G, strat_cls, skill=0.9, detection_rate=0.08)
            stats, _ = sim.run_simulations(n_sims)

            sr = stats['succeeded'] / stats['total'] * 100
            dr = stats['detected'] / stats['total'] * 100
            fr = stats['failed'] / stats['total'] * 100
            avg_time, ci_time, avg_steps, ci_steps = float('nan'), 0, float('nan'), 0
            if stats['succeeded'] > 0:
                all_t = [r['time'] for p in stats['successful_path_details'].values() for r in p]
                all_s = [r['steps'] for p in stats['successful_path_details'].values() for r in p]
                avg_time = np.mean(all_t)
                ci_time = (1.96 * np.std(all_t, ddof=1) / np.sqrt(len(all_t))) if len(all_t) > 1 else 0
                avg_steps = np.mean(all_s)
                ci_steps = (1.96 * np.std(all_s, ddof=1) / np.sqrt(len(all_s))) if len(all_s) > 1 else 0

            # Deviation rate or activations, depending on the strategy
            avg_act = 0
            if hasattr(sim.strategy, 'get_deviation_rate'):
                avg_act = sim.strategy.get_deviation_rate()
            elif hasattr(sim.strategy, '_total_activations'):
                avg_act = sim.strategy._total_activations / max(sim.strategy._total_runs, 1)

            return {
                'type': 'sensitivity', 'run_idx': run_idx,
                'param_key': param_key, 'config_name': config_name,
                'success_rate': sr, 'detection_rate': dr, 'failure_rate': fr,
                'avg_time': avg_time, 'ci_time': ci_time,
                'avg_steps': avg_steps, 'ci_steps': ci_steps,
                'avg_activations': avg_act
            }

    except Exception as e:
        return {
            'type': task_spec['type'], 'run_idx': run_idx,
            'error': f"{type(e).__name__}: {e}\n{traceback.format_exc()}",
            'strat_name': task_spec.get('strat_name', ''),
            'config_name': task_spec.get('config_name', ''),
            'param_key': task_spec.get('param_key', ''),
        }


def serialize_graph(G):
    """Serializes a NetworkX graph so it can be passed to child processes."""
    nodes = []
    for name, attrs in G.nodes(data=True):
        safe_attrs = {}
        for k, v in attrs.items():
            if isinstance(v, (list, dict)):
                safe_attrs[k] = json.dumps(v)
            else:
                safe_attrs[k] = v
        nodes.append((name, safe_attrs))

    edges = []
    for u, v, attrs in G.edges(data=True):
        edges.append((u, v, dict(attrs)))

    return {'nodes': nodes, 'edges': edges}


# Aggregating results

def _empty_accum():
    return defaultdict(lambda: defaultdict(list))


def _accumulate_result(accumulator, config_name, metrics_dict):
    for k, v in metrics_dict.items():
        if k in ('params', 'type', 'run_idx', 'strat_name', 'config_name', 'param_key', 'error'):
            continue
        # Only numeric metrics are aggregated
        if not isinstance(v, (int, float)):
            continue
        accumulator[config_name][k].append(v)
    if 'params' not in accumulator[config_name] and 'params' in metrics_dict:
        accumulator[config_name]['params'] = metrics_dict['params']


def _aggregate(accumulator):
    aggregated = {}
    for config, metrics in accumulator.items():
        agg = {}
        for k, vals in metrics.items():
            if k == 'params':
                agg[k] = vals
                continue
            arr = np.array(vals, dtype=float)
            arr = arr[~np.isnan(arr)]
            if len(arr) == 0:
                agg[k] = float('nan')
                agg[f'{k}_ci'] = 0.0
            else:
                agg[k] = float(np.mean(arr))
                # 95% CI half-width (normal approximation)
                agg[f'{k}_ci'] = float(
                    1.96 * np.std(arr, ddof=1) / np.sqrt(len(arr))
                ) if len(arr) > 1 else 0.0
        aggregated[config] = agg
    return aggregated


def _select_best_strategy(agg_baseline, agg_all_strategies):
    """
    Picks the best strategy versus the baseline, using 95% CIs.
    
    Priority:
    1. Highest success (CI lower bound)
    2. Lowest detection (CI upper bound)
    3. Lowest failure (CI upper bound)
    """
    baseline_name = list(agg_baseline.keys())[0]
    baseline = agg_baseline[baseline_name]
    
    # Baseline CI bounds
    bl_sr_lower = baseline['success_rate'] - baseline.get('success_rate_ci', 0)
    bl_dr_upper = baseline['detection_rate'] + baseline.get('detection_rate_ci', 0)
    bl_fr_upper = baseline['failure_rate'] + baseline.get('failure_rate_ci', 0)
    
    best_strategy = None
    best_score = None
    
    for strat_name, metrics in agg_all_strategies.items():
        # Success lower bound (pessimistic)
        sr_lower = metrics['success_rate'] - metrics.get('success_rate_ci', 0)
        # Detection upper bound (pessimistic)
        dr_upper = metrics['detection_rate'] + metrics.get('detection_rate_ci', 0)
        # Failure upper bound (pessimistic)
        fr_upper = metrics['failure_rate'] + metrics.get('failure_rate_ci', 0)
        
        # Priority tuple
        score = (
            sr_lower,        # 1. highest success
            -dr_upper,       # 2. lowest detection (negated)
            -fr_upper        # 3. lowest failure (negated)
        )
        
        # Tuples compare element by element, so the priorities apply in order
        if best_score is None or score > best_score:
            best_score = score
            best_strategy = strat_name
    
    print(f"\n{'='*70}")
    print(f"SELEZIONE MIGLIORE STRATEGIA (considerando CI95%)")
    print(f"{'='*70}")
    print(f"Baseline ({baseline_name}):")
    print(f"  Successo: {baseline['success_rate']:.1f}% [{bl_sr_lower:.1f}%, {baseline['success_rate'] + baseline.get('success_rate_ci', 0):.1f}%]")
    print(f"  Rilevamento: {baseline['detection_rate']:.1f}% [{baseline['detection_rate'] - baseline.get('detection_rate_ci', 0):.1f}%, {bl_dr_upper:.1f}%]")
    print(f"  Fallimento: {baseline['failure_rate']:.1f}% [{baseline['failure_rate'] - baseline.get('failure_rate_ci', 0):.1f}%, {bl_fr_upper:.1f}%]")
    
    best_metrics = agg_all_strategies[best_strategy]
    sr_lower = best_metrics['success_rate'] - best_metrics.get('success_rate_ci', 0)
    dr_upper = best_metrics['detection_rate'] + best_metrics.get('detection_rate_ci', 0)
    fr_upper = best_metrics['failure_rate'] + best_metrics.get('failure_rate_ci', 0)
    
    print(f"\n✅ Migliore strategia: {best_strategy}")
    print(f"  Successo: {best_metrics['success_rate']:.1f}% [{sr_lower:.1f}%, {best_metrics['success_rate'] + best_metrics.get('success_rate_ci', 0):.1f}%]")
    print(f"  Rilevamento: {best_metrics['detection_rate']:.1f}% [{best_metrics['detection_rate'] - best_metrics.get('detection_rate_ci', 0):.1f}%, {dr_upper:.1f}%]")
    print(f"  Fallimento: {best_metrics['failure_rate']:.1f}% [{best_metrics['failure_rate'] - best_metrics.get('failure_rate_ci', 0):.1f}%, {fr_upper:.1f}%]")
    print(f"{'='*70}\n")
    
    return best_strategy, best_metrics


# Excel export

def export_to_excel(agg_baseline, agg_sensitivity, sensitivity_grid,
                    all_sa_agg, best_name_agg, best_agg,
                    n_runs, n_sims, output_path="risultati_sensitività.xlsx"):
    """Saves all aggregated results to an Excel file, one sheet per topic."""
    if not HAS_OPENPYXL:
        print("⚠️  openpyxl non disponibile, export Excel saltato.")
        return

    wb = openpyxl.Workbook()

    # Styles
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    best_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    # Color per sensitivity parameter
    param_colors = {
        'gamma': 'FFD6D6',           # pastel red
        'T_max': 'FFFFCC',           # pastel yellow
        'L': 'D4E6F1',               # pastel blue
        'delta_T': 'D4F1D4',         # pastel green
        'p_func': 'FFE6CC',          # pastel orange
        'perturbation_p': 'E6CCFF'   # pastel purple (StochasticPerturbation)
    }
    
    def get_param_fill(param_key):
        """Returns the PatternFill for a sensitivity parameter."""
        color = param_colors.get(param_key, 'D3D3D3')  # default grey
        return PatternFill(start_color=color, end_color=color, fill_type="solid")

    def write_header(ws, headers, row=1):
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

    def auto_width(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 3, 40)

    # Sheet: Summary
    ws = wb.active
    ws.title = "Riepilogo"
    ws.cell(row=1, column=1, value="Risultati Aggregati — Analisi di Sensitività SAHybrid").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Data: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    ws.cell(row=3, column=1, value=f"Numero di ripetizioni (N_RUNS): {n_runs}")
    ws.cell(row=4, column=1, value=f"Simulazioni per configurazione per run (N_SIMS): {n_sims}")
    ws.cell(row=6, column=1, value=f"Migliore configurazione: {best_name_agg}").font = Font(bold=True, size=12, color="D4A017")
    ws.cell(row=7, column=1, value=f"  Successo: {best_agg['success_rate']:.1f}% ± {best_agg.get('success_rate_ci',0):.2f}")
    ws.cell(row=8, column=1, value=f"  Rilevamento: {best_agg['detection_rate']:.1f}% ± {best_agg.get('detection_rate_ci',0):.2f}")
    ws.cell(row=9, column=1, value=f"  Tempo medio: {best_agg['avg_time']:.2f}h ± {best_agg.get('avg_time_ci',0):.2f}")
    if 'params' in best_agg:
        ws.cell(row=10, column=1, value=f"  Parametri: {best_agg['params']}")

    # Sheet: Stats and baseline vs config comparison
    ws_stat = wb.create_sheet("Confronto Statistico")
    ws_stat.cell(row=1, column=1, value="ANALISI STATISTICA: Confronto Baseline vs Configurazioni").font = Font(bold=True, size=12)
    ws_stat.cell(row=2, column=1, value=f"Data: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    
    ws_stat.cell(row=4, column=1, value="BASELINE (MaxProbability):").font = Font(bold=True, size=11)
    baseline_maxprob = agg_baseline.get('MaxProbabilityStrategy', {})
    if not baseline_maxprob and agg_baseline:
        baseline_maxprob = list(agg_baseline.values())[0]
    
    baseline_sr = baseline_maxprob.get('success_rate', 0)
    baseline_sr_ci = baseline_maxprob.get('success_rate_ci', 0)
    baseline_dr = baseline_maxprob.get('detection_rate', 0)
    baseline_dr_ci = baseline_maxprob.get('detection_rate_ci', 0)
    baseline_time = baseline_maxprob.get('avg_time', 0)
    baseline_time_ci = baseline_maxprob.get('avg_time_ci', 0)
    
    ws_stat.cell(row=5, column=1, value=f"  Successo: {baseline_sr:.1f}% ± {baseline_sr_ci:.2f} (95% CI)")
    ws_stat.cell(row=6, column=1, value=f"  Rilevamento: {baseline_dr:.1f}% ± {baseline_dr_ci:.2f} (95% CI)")
    ws_stat.cell(row=7, column=1, value=f"  Tempo Medio: {baseline_time:.2f}h ± {baseline_time_ci:.2f}")
    
    ws_stat.cell(row=9, column=1, value="INTERPRETAZIONE INTERVALLI DI CONFIDENZA:").font = Font(bold=True, size=11)
    ws_stat.cell(row=10, column=1, value="  • Se IC95% NON si sovrappongono → il miglioramento è STATISTICAMENTE SIGNIFICATIVO")
    ws_stat.cell(row=11, column=1, value="  • Se IC95% SI sovrappongono → il cambiamento NON è significativo")
    ws_stat.cell(row=12, column=1, value="  • Confronto: se uno strisce è verde/rosso → c'è miglioramento/peggioramento statisticamente significativo")
    
    ws_stat.cell(row=14, column=1, value="OBIETTIVO:").font = Font(bold=True, size=11)
    ws_stat.cell(row=15, column=1, value="  • Ridurre Detection Rate (DR) possibilmente migliorando/mantenendo Success Rate (SR)")
    ws_stat.cell(row=16, column=1, value="  • Strategia: Disturbare MaxProb con deviazioni probabilistiche per ridurne la prevedibilità")
    
    # Comparison table, one row per config
    ws_stat.cell(row=18, column=1, value="CONFIG").font = Font(bold=True)
    ws_stat.cell(row=18, column=2, value="Success ±CI").font = Font(bold=True)
    ws_stat.cell(row=18, column=3, value="vs Baseline").font = Font(bold=True)
    ws_stat.cell(row=18, column=4, value="Detection ±CI").font = Font(bold=True)
    ws_stat.cell(row=18, column=5, value="vs Baseline").font = Font(bold=True)
    ws_stat.cell(row=18, column=6, value="Conclusione").font = Font(bold=True)
    
    def intervals_overlap(val1, ci1, val2, ci2):
        low1, high1 = val1 - ci1, val1 + ci1
        low2, high2 = val2 - ci2, val2 + ci2
        return not (high1 < low2 or high2 < low1)
    
    row_idx = 19
    if all_sa_agg:
        for name in sorted(all_sa_agg.keys()):
            m = all_sa_agg[name]
            if 'perturbation_p' not in name:
                continue
            
            sr = m.get('success_rate', float('nan'))
            sr_ci = m.get('success_rate_ci', 0)
            dr = m.get('detection_rate', float('nan'))
            dr_ci = m.get('detection_rate_ci', 0)
            
            # Do the confidence intervals overlap?
            sr_overlap = intervals_overlap(sr, sr_ci, baseline_sr, baseline_sr_ci)
            dr_overlap = intervals_overlap(dr, dr_ci, baseline_dr, baseline_dr_ci)
            
            sr_cond = "UGUALE" if sr_overlap else ("MIGLIORA" if sr > baseline_sr else "PEGGIORA")
            dr_cond = "UGUALE" if dr_overlap else ("MIGLIORA" if dr < baseline_dr else "PEGGIORA")
            
            if sr_cond == "MIGLIORA" and dr_cond == "MIGLIORA":
                conclusion = "✓ MIGLIORAMENTO DOPPIO"
                color = "C6EFCE"
            elif sr_cond == "PEGGIORA" or dr_cond == "PEGGIORA":
                conclusion = "✗ ALMENO UNO PEGGIORA"
                color = "FFC7CE"
            else:
                conclusion = "~ NESSUN CAMB. SIGN."
                color = "FFFFCC"
            
            ws_stat.cell(row=row_idx, column=1, value=name)
            ws_stat.cell(row=row_idx, column=2, value=f"{sr:.1f}% ± {sr_ci:.2f}")
            ws_stat.cell(row=row_idx, column=3, value=sr_cond)
            ws_stat.cell(row=row_idx, column=4, value=f"{dr:.1f}% ± {dr_ci:.2f}")
            ws_stat.cell(row=row_idx, column=5, value=dr_cond)
            cell_c = ws_stat.cell(row=row_idx, column=6, value=conclusion)
            cell_c.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            
            row_idx += 1
    
    # Sheet: Baseline
    ws_bl = wb.create_sheet("Baseline")
    bl_headers = ["Strategia", "Successo (%)", "± CI95", "Rilevamento (%)", "± CI95",
                  "Fallimento (%)", "± CI95", "Tempo Medio (h)", "± CI95",
                  "Passi Medi", "± CI95"]
    write_header(ws_bl, bl_headers)
    for row_idx, (name, m) in enumerate(agg_baseline.items(), 2):
        vals = [
            name,
            round(m.get('success_rate', float('nan')), 2),
            round(m.get('success_rate_ci', 0), 2),
            round(m.get('detection_rate', float('nan')), 2),
            round(m.get('detection_rate_ci', 0), 2),
            round(m.get('failure_rate', float('nan')), 2),
            round(m.get('failure_rate_ci', 0), 2),
            round(m.get('avg_time', float('nan')), 2),
            round(m.get('avg_time_ci', 0), 2),
            round(m.get('avg_steps', float('nan')), 2),
            round(m.get('avg_steps_ci', 0), 2),
        ]
        for col_idx, v in enumerate(vals, 1):
            cell = ws_bl.cell(row=row_idx, column=col_idx, value=v)
            cell.border = thin_border
    auto_width(ws_bl)

    # Sheet: Sensitivity by parameter
    for key, (param_name, _) in sensitivity_grid.items():
        ws_p = wb.create_sheet(f"Sens_{param_name}")
        
        # Columns depend on the parameter
        if param_name == 'perturbation_p':
            p_headers = [
                "Configurazione", "Successo (%)", "± CI95", "Rilevamento (%)", "± CI95",
                "Fallimento (%)", "± CI95", "Tempo Medio (h)", "± CI95",
                "Passi Medi", "± CI95", "Deviation Rate (%)", "± CI95"
            ]
        else:
            p_headers = [
                "Configurazione", "Successo (%)", "± CI95", "Rilevamento (%)", "± CI95",
                "Fallimento (%)", "± CI95", "Tempo Medio (h)", "± CI95",
                "Passi Medi", "± CI95", "MaxProb Act/run", "± CI95"
            ]
        
        write_header(ws_p, p_headers)
        
        # Use this parameter's color
        param_fill = get_param_fill(key)
        
        for row_idx, (cfg, m) in enumerate(agg_sensitivity[key].items(), 2):
            vals = [
                cfg,
                round(m.get('success_rate', float('nan')), 2),
                round(m.get('success_rate_ci', 0), 2),
                round(m.get('detection_rate', float('nan')), 2),
                round(m.get('detection_rate_ci', 0), 2),
                round(m.get('failure_rate', float('nan')), 2),
                round(m.get('failure_rate_ci', 0), 2),
                round(m.get('avg_time', float('nan')), 2),
                round(m.get('avg_time_ci', 0), 2),
                round(m.get('avg_steps', float('nan')), 2),
                round(m.get('avg_steps_ci', 0), 2),
                round(m.get('avg_activations', float('nan')), 2),
                round(m.get('avg_activations_ci', 0), 2),
            ]
            for col_idx, v in enumerate(vals, 1):
                cell = ws_p.cell(row=row_idx, column=col_idx, value=v)
                cell.border = thin_border
                if param_fill:
                    cell.fill = param_fill
        auto_width(ws_p)

    # Sheet: All configs
    ws_all = wb.create_sheet("Tutte le Config")
    all_headers = [
        "Configurazione", "Successo (%)", "± CI95", "Rilevamento (%)", "± CI95",
        "Fallimento (%)", "± CI95", "Tempo Medio (h)", "± CI95",
        "Passi Medi", "± CI95", "MaxProb Act/run", "± CI95",
        "Score (Succ - Rilev)"
    ]
    write_header(ws_all, all_headers)

    # Baseline row
    row_idx = 2
    for name, m in agg_baseline.items():
        score = m.get('success_rate', 0) - m.get('detection_rate', 0)
        vals = [
            name,
            round(m.get('success_rate', float('nan')), 2),
            round(m.get('success_rate_ci', 0), 2),
            round(m.get('detection_rate', float('nan')), 2),
            round(m.get('detection_rate_ci', 0), 2),
            round(m.get('failure_rate', float('nan')), 2),
            round(m.get('failure_rate_ci', 0), 2),
            round(m.get('avg_time', float('nan')), 2),
            round(m.get('avg_time_ci', 0), 2),
            round(m.get('avg_steps', float('nan')), 2),
            round(m.get('avg_steps_ci', 0), 2),
            0, 0,
            round(score, 2),
        ]
        for col_idx, v in enumerate(vals, 1):
            cell = ws_all.cell(row=row_idx, column=col_idx, value=v)
            cell.border = thin_border
        row_idx += 1

    # All SA configs
    for label, m in all_sa_agg.items():
        score = m.get('success_rate', 0) - m.get('detection_rate', 0)
        is_best = (label == best_name_agg)
        vals = [
            label,
            round(m.get('success_rate', float('nan')), 2),
            round(m.get('success_rate_ci', 0), 2),
            round(m.get('detection_rate', float('nan')), 2),
            round(m.get('detection_rate_ci', 0), 2),
            round(m.get('failure_rate', float('nan')), 2),
            round(m.get('failure_rate_ci', 0), 2),
            round(m.get('avg_time', float('nan')), 2),
            round(m.get('avg_time_ci', 0), 2),
            round(m.get('avg_steps', float('nan')), 2),
            round(m.get('avg_steps_ci', 0), 2),
            round(m.get('avg_activations', float('nan')), 2),
            round(m.get('avg_activations_ci', 0), 2),
            round(score, 2),
        ]
        for col_idx, v in enumerate(vals, 1):
            cell = ws_all.cell(row=row_idx, column=col_idx, value=v)
            cell.border = thin_border
            if is_best:
                cell.fill = best_fill
        row_idx += 1
    auto_width(ws_all)

    # Sheet: Raw data per run
    # Added from main once raw_results are collected
    # (created by the caller)

    wb.save(output_path)
    print(f"📗 Excel salvato: {output_path}")
    return wb


def add_raw_results_sheet(wb, all_raw_results, output_path, param_colors=None):
    """Adds a sheet with the raw results of every run, sorted by sensitivity and run."""
    ws_raw = wb.create_sheet("Dati Grezzi (per run)")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    
    if param_colors is None:
        param_colors = {}
    
    def get_param_fill(param_key):
        """Returns the PatternFill for a sensitivity parameter."""
        if param_key == '-' or param_key == 'baseline':
            return None
        color = param_colors.get(param_key, 'D3D3D3')  # default grey
        return PatternFill(start_color=color, end_color=color, fill_type="solid")

    raw_headers = [
        "Run", "Tipo", "Strategia/Config", "Parametro Variato",
        "Successo (%)", "Rilevamento (%)", "Fallimento (%)",
        "Tempo Medio (h)", "CI Tempo", "Passi Medi", "CI Passi",
        "MaxProb Act/run"
    ]
    for col_idx, h in enumerate(raw_headers, 1):
        cell = ws_raw.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    # Sort: baseline first, then by param_key, then by run_idx
    def sort_key(result):
        tipo = result.get('type', '')
        param_key = result.get('param_key', '-')
        run_idx = result.get('run_idx', 0)
        
        # Baseline sorts first
        tipo_priority = 0 if tipo == 'baseline' else 1
        return (tipo_priority, param_key, run_idx)
    
    sorted_results = sorted(all_raw_results, key=sort_key)
    
    for row_idx, result in enumerate(sorted_results, 2):
        param_key = result.get('param_key', '-')
        param_fill = get_param_fill(param_key)
        
        if 'error' in result:
            vals = [
                result.get('run_idx', '?'), result.get('type', '?'),
                result.get('strat_name', result.get('config_name', '?')),
                result.get('param_key', '-'),
                "ERRORE", "", "", "", "", "", "", ""
            ]
        elif result['type'] == 'baseline':
            vals = [
                result['run_idx'], 'baseline', result['strat_name'], '-',
                round(result['success_rate'], 2), round(result['detection_rate'], 2),
                round(result['failure_rate'], 2),
                round(result['avg_time'], 2), round(result['ci_time'], 2),
                round(result['avg_steps'], 2), round(result['ci_steps'], 2),
                0
            ]
        else:
            vals = [
                result['run_idx'], 'sensitivity', result['config_name'],
                result.get('param_key', '-'),
                round(result['success_rate'], 2), round(result['detection_rate'], 2),
                round(result['failure_rate'], 2),
                round(result['avg_time'], 2), round(result['ci_time'], 2),
                round(result['avg_steps'], 2), round(result['ci_steps'], 2),
                round(result.get('avg_activations', 0), 2)
            ]
        
        for col_idx, v in enumerate(vals, 1):
            cell = ws_raw.cell(row=row_idx, column=col_idx, value=v)
            cell.border = thin_border
            if param_fill:
                cell.fill = param_fill

    # Auto-fit column widths
    for col in ws_raw.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws_raw.column_dimensions[col_letter].width = min(max_len + 3, 40)

    wb.save(output_path)
    print(f"📗 Foglio 'Dati Grezzi' aggiunto a: {output_path}")





# Main (parallel run)

def main():
    nest_asyncio.apply()

    # Config
    NVD_API_KEY = None
    CSV_FILE_PATH = 'glpi.csv'
    N_RIPETIZIONI = 100    # Repetitions per config
    N_SIMS = 200000       # Simulations per config, per repetition
    MAX_WORKERS = None   # None = all available CPUs
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    EXCEL_OUTPUT = f"risultati_sensitività_{timestamp}.xlsx"
    TXT_OUTPUT = "risultati_sensitività.txt"

    # StochasticPerturbation experiment: p from 0.5% to 10%
    # (converted from percent to fraction below)
    p_values_percent = [0.5, 1.0, 1.5, 2.0, 2.5, 3.0, 4.0, 5.0, 7.0, 10.0]
    p_values = [p / 100.0 for p in p_values_percent]  # Percent to fraction
    
    # Grid for the new experiment (comment out the previous one)
    sensitivity_grid = {
        'perturbation_p': ('perturbation_p', p_values),  # Deviation probability
    }
    


    # Load data
    print("=" * 70)
    print("  CARICAMENTO DIGITAL TWIN E ARRICCHIMENTO CVE")
    print("=" * 70)

    dt = DigitalTwin()
    try:
        dt.load_from_csv(CSV_FILE_PATH)
    except FileNotFoundError:
        print(f"ERRORE: file '{CSV_FILE_PATH}' non trovato.")
        return

    enricher = CVEEnricher(dt.get_graph(), api_key=NVD_API_KEY)
    asyncio.run(enricher.run_enrichment())

    graph = dt.get_graph()
    graph_data = serialize_graph(graph)

    # Build tasks
    print("\n" + "=" * 70)
    print("  CREAZIONE TASK PER ESECUZIONE PARALLELA")
    print("=" * 70)

    tasks = []
    rep_idx_counter = 0
    
    # Phase 1: MaxProb baseline (MaxProb only, for a direct comparison)
    print("\n  ▶ Creando task BASELINE (MaxProbabilityStrategy)...")
    strat_name = 'MaxProb (baseline)'
    for rep_idx in range(1, N_RIPETIZIONI + 1):
        rep_idx_counter += 1
        tasks.append((graph_data, {
            'type': 'baseline', 'run_idx': rep_idx, 'run_total': N_RIPETIZIONI,
            'strat_name': strat_name, 'n_sims': N_SIMS
        }))
    print(f"     + {N_RIPETIZIONI} run baseline MaxProb")
    
    # Phase 2: StochasticPerturbation at different p values
    print(f"\n  ▶ Creando task SENSITIVITÀ (StochasticPerturbationStrategy)...")
    for key, (param_name, values) in sensitivity_grid.items():
        for val in values:
            # Pass p directly
            config_name = f"p={val:.3f}" if isinstance(val, float) else f"p={val}"
            for rep_idx in range(1, N_RIPETIZIONI + 1):
                rep_idx_counter += 1
                tasks.append((graph_data, {
                    'type': 'sensitivity', 'run_idx': rep_idx, 'run_total': N_RIPETIZIONI,
                    'param_key': key, 'config_name': config_name,
                    'stochastic_p': val, 'deviation_mode': 'stealth',
                    'n_sims': N_SIMS
                }))
        print(f"     + {len(values)} configurazioni × {N_RIPETIZIONI} run ciascuna")

    total_tasks = len(tasks)
    total_sims = total_tasks * N_SIMS
    n_workers = MAX_WORKERS or min(multiprocessing.cpu_count(), total_tasks)

    print(f"  Task totali: {total_tasks}")
    print(f"  Simulazioni totali: {total_sims:,}")
    print(f"  Worker paralleli: {n_workers}")
    print(f"  N_RIPETIZIONI per configurazione: {N_RIPETIZIONI}  |  N_SIMS per ripetizione: {N_SIMS}")

    # Run in parallel
    print("\n" + "=" * 70)
    print("  ESECUZIONE PARALLELA IN CORSO...")
    print("=" * 70)

    all_raw_results = []
    completed = 0
    errors = 0
    start_time = datetime.now()

    with ProcessPoolExecutor(max_workers=n_workers) as executor:
        futures = {executor.submit(_worker_run_simulation, task): task for task in tasks}

        # Collect results as workers finish
        for future in as_completed(futures):
            completed += 1
            result = future.result()
            all_raw_results.append(result)

            if 'error' in result:
                errors += 1
                print(f"  [{completed}/{total_tasks}] ❌ ERRORE: {result.get('config_name', result.get('strat_name', '?'))}")
            else:
                label = result.get('strat_name', result.get('config_name', '?'))
                rep_idx = result.get('run_idx', '?')
                rep_total = result.get('run_total', '?')
                elapsed = (datetime.now() - start_time).total_seconds()
                # ETA from the average time per task so far
                eta = (elapsed / completed) * (total_tasks - completed)
                print(f"  [{completed}/{total_tasks}] ✅ Rep {rep_idx}/{rep_total} | {label:.<35} "
                      f"Succ: {result['success_rate']:5.1f}% | "
                      f"ETA: {eta/60:.1f} min")

    elapsed_total = (datetime.now() - start_time).total_seconds()
    print(f"\n⏱️  Tempo totale: {elapsed_total/60:.1f} minuti ({elapsed_total:.0f}s)")
    print(f"  Completati: {completed - errors}/{total_tasks}  |  Errori: {errors}")

    # Aggregate
    print("\n" + "=" * 70)
    print("  AGGREGAZIONE RISULTATI")
    print("=" * 70)

    baseline_acc = _empty_accum()
    sensitivity_acc = {key: _empty_accum() for key in sensitivity_grid}

    for result in all_raw_results:
        if 'error' in result:
            continue
        if result['type'] == 'baseline':
            _accumulate_result(baseline_acc, result['strat_name'], result)
        elif result['type'] == 'sensitivity':
            _accumulate_result(sensitivity_acc[result['param_key']], result['config_name'], result)

    agg_baseline = _aggregate(baseline_acc)
    agg_sensitivity = {}
    for key in sensitivity_grid:
        agg_sensitivity[key] = _aggregate(sensitivity_acc[key])

    # All aggregated SA configs
    all_sa_agg = {}
    for key, res in agg_sensitivity.items():
        for cfg, m in res.items():
            all_sa_agg[f"{key}:{cfg}"] = m

    # Best config (using 95% CIs)
    if all_sa_agg:
        best_name_agg, best_agg = _select_best_strategy(agg_baseline, all_sa_agg)
    else:
        best_name_agg = "N/A"
        best_agg = {}

    # Print results
    print("\n" + "=" * 100)
    print(f"📊 RISULTATI AGGREGATI SU {N_RIPETIZIONI} RIPETIZIONI")
    print("=" * 100)

    print("\n--- Baseline ---")
    for name, m in agg_baseline.items():
        print(f"  {name:.<30} Succ: {m['success_rate']:.1f}% ± {m.get('success_rate_ci',0):.2f} | "
              f"Rilev: {m['detection_rate']:.1f}% ± {m.get('detection_rate_ci',0):.2f} | "
              f"Tempo: {m['avg_time']:.2f}h ± {m.get('avg_time_ci',0):.2f}")

    for key, (param_name, _) in sensitivity_grid.items():
        print(f"\n--- Sensitività: {param_name} ---")
        for cfg, m in agg_sensitivity[key].items():
            print(f"  {cfg:.<30} Succ: {m['success_rate']:.1f}% ± {m.get('success_rate_ci',0):.2f} | "
                  f"Rilev: {m['detection_rate']:.1f}% ± {m.get('detection_rate_ci',0):.2f} | "
                  f"Tempo: {m['avg_time']:.2f}h ± {m.get('avg_time_ci',0):.2f} | "
                  f"MaxP/rep: {m.get('avg_activations',0):.2f} ± {m.get('avg_activations_ci',0):.2f}")

    print(f"\n{'=' * 100}")
    print(f"MIGLIORE CONFIGURAZIONE: {best_name_agg}")
    if best_agg:
        print(f"  Parametri: {best_agg.get('params', 'N/A')}")
        print(f"  Successo:    {best_agg['success_rate']:.1f}% ± {best_agg.get('success_rate_ci',0):.2f}")
        print(f"  Rilevamento: {best_agg['detection_rate']:.1f}% ± {best_agg.get('detection_rate_ci',0):.2f}")
        print(f"  Tempo medio: {best_agg['avg_time']:.2f}h ± {best_agg.get('avg_time_ci',0):.2f}")
    print("=" * 100)

    # Save TXT
    lines = []
    lines.append("=" * 100)
    lines.append(f"RISULTATI AGGREGATI SU {N_RIPETIZIONI} RIPETIZIONI PER CONFIGURAZIONE — {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append(f"Simulazioni per configurazione per ripetizione: {N_SIMS}")
    lines.append(f"Tempo totale esecuzione: {elapsed_total/60:.1f} minuti")
    lines.append("=" * 100)
    lines.append("\n--- Baseline ---")
    for name, m in agg_baseline.items():
        lines.append(f"  {name:.<30} Succ: {m['success_rate']:.1f}% ± {m.get('success_rate_ci',0):.2f} | "
                     f"Rilev: {m['detection_rate']:.1f}% ± {m.get('detection_rate_ci',0):.2f} | "
                     f"Tempo: {m['avg_time']:.2f}h ± {m.get('avg_time_ci',0):.2f}")
    for key, (param_name, _) in sensitivity_grid.items():
        lines.append(f"\n--- Sensitività: {param_name} ---")
        for cfg, m in agg_sensitivity[key].items():
            lines.append(f"  {cfg:.<30} Succ: {m['success_rate']:.1f}% ± {m.get('success_rate_ci',0):.2f} | "
                         f"Rilev: {m['detection_rate']:.1f}% ± {m.get('detection_rate_ci',0):.2f} | "
                         f"Tempo: {m['avg_time']:.2f}h ± {m.get('avg_time_ci',0):.2f} | "
                         f"MaxP/run: {m.get('avg_activations',0):.2f} ± {m.get('avg_activations_ci',0):.2f}")
    lines.append(f"\n{'=' * 100}")
    lines.append(f"MIGLIORE: {best_name_agg}")
    if best_agg:
        lines.append(f"  Parametri: {best_agg.get('params', 'N/A')}")
        lines.append(f"  Successo: {best_agg['success_rate']:.1f}% ± {best_agg.get('success_rate_ci',0):.2f}")
        lines.append(f"  Rilevamento: {best_agg['detection_rate']:.1f}% ± {best_agg.get('detection_rate_ci',0):.2f}")
    lines.append("=" * 100)
    with open(TXT_OUTPUT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print(f"\n📄 TXT salvato: {TXT_OUTPUT}")

    # Export Excel
    print("\n" + "=" * 70)
    print("  EXPORT EXCEL")
    print("=" * 70)
    wb = export_to_excel(
        agg_baseline, agg_sensitivity, sensitivity_grid,
        all_sa_agg, best_name_agg, best_agg,
        N_RIPETIZIONI, N_SIMS, output_path=EXCEL_OUTPUT
    )
    if wb:
        param_colors = {
            'gamma': 'FFD6D6',      # pastel red
            'T_max': 'FFFFCC',      # pastel yellow
            'L': 'D4E6F1',          # pastel blue
            'delta_T': 'D4F1D4',    # pastel green
            'p_func': 'FFE6CC'      # pastel orange
        }
        wb = openpyxl.load_workbook(EXCEL_OUTPUT)
        add_raw_results_sheet(wb, all_raw_results, EXCEL_OUTPUT, param_colors)

    print("\n" + "=" * 70)
    print("  ✅ ESECUZIONE COMPLETATA")
    print("=" * 70)


if __name__ == "__main__":
    main()